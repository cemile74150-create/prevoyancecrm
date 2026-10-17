"""Calcul LAMal à partir des fichiers officiels OFSP / EDI / swisstopo.

Les montants ne sont jamais inventés : une offre n'existe que si une ligne
officielle correspond exactement au profil (canton, région, âge, sous-groupe,
accident, franchise, tarif).
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Iterable, Optional

CENT = Decimal("0.01")
TWELVE = Decimal(12)
ADULT_COINSURANCE_CAP = Decimal("700")
CHILD_COINSURANCE_CAP = Decimal("350")

# Cantons à région unique et cantons subdivisés, tels que publiés par l'OFSP
# pour 2027 (ordonnance du DFI, état au 1.1.2027) et confirmés par Prämien_CH.
SINGLE_REGION_CANTONS = frozenset({
    "AG", "AI", "AR", "BS", "GE", "GL", "JU", "NE", "NW", "OW", "SO", "SZ", "TG", "UR", "ZG",
})
MULTI_REGION_CANTONS = {
    "BE": {1, 2, 3},
    "BL": {1, 2},
    "FR": {1, 2},
    "GR": {1, 2, 3},
    "LU": {1, 2, 3},
    "SG": {1, 2, 3},
    "SH": {1, 2},
    "TI": {1, 2},
    "VD": {1, 2},
    "VS": {1, 2},
    "ZH": {1, 2, 3},
}
OFFICIAL_CANTONS = SINGLE_REGION_CANTONS | frozenset(MULTI_REGION_CANTONS)

AGE_CLASS_FR = {
    "AKA_01_KIN": "Enfant (0–18 ans)",
    "AKA_02_JUG": "Jeune adulte (19–25 ans)",
    "AKA_03_ERW": "Adulte (dès 26 ans)",
}
ORDINARY_SUBGROUP = {
    "AKA_01_KIN": "K1",
    "AKA_02_JUG": "J1",
    "AKA_03_ERW": "E1",
}
TARIFTYP_FR = {
    "BASE": "Modèle standard",
    "PRAXIS": "Médecin de famille",
    "TEL_DIG": "Télémédecine",
    "FLEX": "Modèle flexible",
    "PHARM": "Pharmacie",
}
ACCIDENT_WITH = "MIT_UNF"
ACCIDENT_WITHOUT = "OHN_UNF"
TERRITORY_CH = "P_OKPCH"
MIN_PREMIUM_ROWS = 100_000

_FRANCHISE_RE = re.compile(r"^FRA_\d+_([EJK])_(\d+)$")
_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
# Exonymes officiels français quand l'annexe EDI est en allemand.
COMMUNE_ALIASES = {
    "berne": "bern",
    "genf": "geneve",
    "geneva": "geneve",
    "lucerne": "luzern",
    "basel": "bale",
    "bale": "basel",
}

SOURCES = {
    "premiums": "https://opendata.swiss/fr/dataset/health-insurance-premiums",
    "premiums_download": "https://opendata.bagnet.ch/?r=/download&path=L1ByYWVtaWVuL1Byw6RtaWVuX0NILmNzdg%3D%3D",
    "regions": "https://www.bag.admin.ch/de/krankenversicherung-praemienregionen",
    "regions_download": "https://www.bag.admin.ch/dam/de/sd-web/Xm8w51dMKXLP/Pr%C3%A4mienregionen%20(g%C3%BCltig%20ab%201.1.2027).xlsx",
    "localities": "https://data.geo.admin.ch/ch.swisstopo-vd.ortschaftenverzeichnis_plz/ortschaftenverzeichnis_plz/ortschaftenverzeichnis_plz_2056.csv.zip",
    "published": "2026-09-29",
    "regions_dated": "2026-09-09",
    "ordinance": "https://www.fedlex.admin.ch/eli/oc/2026/483/fr",
}


def money(value) -> Decimal:
    return Decimal(str(value).strip()).quantize(CENT, rounding=ROUND_HALF_UP)


def money_str(value: Decimal) -> str:
    return f"{money(value):.2f}"


def annual_from_month(month: Decimal) -> Decimal:
    return money(month * TWELVE)


def fold_name(value: Optional[str]) -> str:
    text = unicodedata.normalize("NFKD", value or "")
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.casefold().replace("’", "'").replace("œ", "oe")
    text = text.replace("-", " ").replace("'", " ").replace("/", " ")
    text = re.sub(r"\([^)]*\)", " ", text)
    text = re.sub(r"[^a-z0-9 ]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return COMMUNE_ALIASES.get(text, text)


def age_class_for_birth_year(birth_year: int, tariff_year: int) -> str:
    """Âge atteint dans l'année de primes : année tarifaire − année de naissance."""
    age = tariff_year - int(birth_year)
    if age < 0 or age > 120:
        raise ValueError("Année de naissance incompatible avec l'année de primes")
    if age <= 18:
        return "AKA_01_KIN"
    if age <= 25:
        return "AKA_02_JUG"
    return "AKA_03_ERW"


def parse_franchise_code(code: str) -> tuple[str, int]:
    match = _FRANCHISE_RE.match((code or "").strip())
    if not match:
        raise ValueError(f"Code de franchise inconnu : {code}")
    return match.group(1), int(match.group(2))


def accident_code(with_accident: bool) -> str:
    return ACCIDENT_WITH if with_accident else ACCIDENT_WITHOUT


def region_code_for(canton: str, region_number: Optional[int]) -> str:
    canton = (canton or "").strip().upper()
    if canton in SINGLE_REGION_CANTONS:
        return "PR_REG_0"
    if region_number is None:
        raise ValueError(f"Région de primes manquante pour {canton}")
    return f"PR_REG_{int(region_number)}"


def coinsurance_cap(age_class: str) -> Decimal:
    if age_class == "AKA_01_KIN":
        return CHILD_COINSURANCE_CAP
    return ADULT_COINSURANCE_CAP


def participation_ceiling(franchise_chf: int, age_class: str) -> Decimal:
    return money(Decimal(franchise_chf) + coinsurance_cap(age_class))


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _open_text(path: Path):
    return path.open("r", encoding="utf-8-sig", newline="")


@dataclass
class ScanReport:
    ok: bool
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    row_count: int = 0
    offer_count: int = 0
    year: Optional[int] = None
    cantons: set[str] = field(default_factory=set)
    quarantined_cantons: dict[str, int] = field(default_factory=dict)
    sha256: str = ""

    def reject(self, message: str):
        self.errors.append(message)
        self.ok = False


@dataclass
class Offer:
    insurer_id: str
    canton: str
    region: str
    age_class: str
    age_subgroup: str
    accident: str
    franchise_code: str
    franchise_chf: int
    ordinary_franchise: bool
    tarif_id: str
    tariftyp: str
    premium_month: Decimal
    base_tariff: bool
    tarif_name_de: str
    tarif_name_fr: str = ""
    insurer_name: str = ""

    @property
    def key(self) -> tuple:
        return (
            self.insurer_id, self.canton, self.region, self.age_class, self.age_subgroup,
            self.accident, self.franchise_code, self.tarif_id,
        )

    def public(self) -> dict:
        month = money(self.premium_month)
        name = self.insurer_name or f"Assureur OFSP n° {self.insurer_id}"
        model = self.tarif_name_fr or TARIFTYP_FR.get(self.tariftyp, self.tariftyp)
        return {
            "insurer_id": self.insurer_id,
            "insurer_name": name,
            "model_type": self.tariftyp,
            "model_label": TARIFTYP_FR.get(self.tariftyp, self.tariftyp),
            "product_name": model,
            "tarif_id": self.tarif_id,
            "franchise_chf": self.franchise_chf,
            "ordinary_franchise": self.ordinary_franchise,
            "accident": self.accident,
            "accident_label": "Avec accidents" if self.accident == ACCIDENT_WITH else "Sans accidents",
            "premium_month": money_str(month),
            "premium_year": money_str(annual_from_month(month)),
            "age_class": self.age_class,
            "age_subgroup": self.age_subgroup,
            "region": self.region,
            "canton": self.canton,
            "base_tariff": self.base_tariff,
            "participation_ceiling_chf": money_str(participation_ceiling(self.franchise_chf, self.age_class)),
        }


def _header_map(fieldnames: list[str]) -> dict[str, str]:
    """Retrouve les colonnes malgré les variantes d'encodage du libellé."""
    wanted = {
        "versicherer": "insurer",
        "kanton": "canton",
        "hoheitsgebiet": "territory",
        "geschaftsjahr": "year",
        "geschaeftsjahr": "year",
        "erhebungsjahr": "survey_year",
        "region": "region",
        "altersklasse": "age_class",
        "unfalleinschluss": "accident",
        "tarif": "tarif",
        "tariftyp": "tariftyp",
        "altersuntergruppe": "subgroup",
        "franchise": "franchise",
        "pramie": "premium",
        "isbasep": "base_p",
        "isbasef": "base_f",
        "tarifbezeichnung": "name_de",
        "eingeschrankt": "restricted",
        "gemeindenbfs": "bfs_list",
        "namefr": "name_fr",
        "kategorie": "category",
    }
    mapping = {}
    for original in fieldnames or []:
        folded = fold_name(original).replace(" ", "")
        if folded in wanted:
            mapping[wanted[folded]] = original
    return mapping


def scan_premium_csv(path: Path, *, tariff_year: int = 2027) -> tuple[ScanReport, dict[tuple, Offer]]:
    report = ScanReport(ok=True, sha256=sha256_file(path), year=tariff_year)
    offers: dict[tuple, Offer] = {}
    with _open_text(path) as handle:
        reader = csv.DictReader(handle)
        columns = _header_map(reader.fieldnames or [])
        required = ["insurer", "canton", "territory", "year", "region", "age_class", "accident", "tarif", "tariftyp", "subgroup", "franchise", "premium"]
        missing = [name for name in required if name not in columns]
        if missing:
            report.reject("Colonnes officielles manquantes dans les primes : " + ", ".join(missing))
            return report, {}
        for row in reader:
            report.row_count += 1
            try:
                year = int(row[columns["year"]])
                canton = (row[columns["canton"]] or "").strip().upper()
                territory = (row[columns["territory"]] or "").strip()
                premium = money(row[columns["premium"]])
                franchise_group, franchise_chf = parse_franchise_code(row[columns["franchise"]])
            except Exception as exc:
                report.reject(f"Ligne {report.row_count} illisible : {exc}")
                if len(report.errors) > 20:
                    break
                continue
            if year != tariff_year:
                report.reject(f"Année de primes {year} au lieu de {tariff_year}")
                break
            if territory != TERRITORY_CH:
                report.reject(f"Territoire non suisse exclu du comparateur : {territory}")
                break
            if premium <= 0:
                report.reject("Prime nulle ou négative")
                break
            age_class = row[columns["age_class"]].strip()
            if age_class not in AGE_CLASS_FR:
                report.reject(f"Classe d'âge inconnue : {age_class}")
                break
            accident = row[columns["accident"]].strip()
            if accident not in {ACCIDENT_WITH, ACCIDENT_WITHOUT}:
                report.reject(f"Couverture accidents inconnue : {accident}")
                break
            if canton not in OFFICIAL_CANTONS:
                report.quarantined_cantons[canton] = report.quarantined_cantons.get(canton, 0) + 1
                continue
            offer = Offer(
                insurer_id=str(row[columns["insurer"]]).strip(),
                canton=canton,
                region=row[columns["region"]].strip(),
                age_class=age_class,
                age_subgroup=row[columns["subgroup"]].strip(),
                accident=accident,
                franchise_code=row[columns["franchise"]].strip(),
                franchise_chf=franchise_chf,
                ordinary_franchise=(row.get(columns.get("base_f", ""), "0") == "1"),
                tarif_id=row[columns["tarif"]].strip(),
                tariftyp=row[columns["tariftyp"]].strip(),
                premium_month=premium,
                base_tariff=(row.get(columns.get("base_p", ""), "0") == "1"),
                tarif_name_de=(row.get(columns.get("name_de", ""), "") or "").strip(),
            )
            previous = offers.get(offer.key)
            if previous and previous.premium_month != offer.premium_month:
                report.reject(f"Prime contradictoire pour {offer.key}")
                break
            offers[offer.key] = offer
            report.cantons.add(canton)
            _ = franchise_group
    report.offer_count = len(offers)
    if report.row_count < MIN_PREMIUM_ROWS and path.stat().st_size > 1_000_000:
        report.reject(f"Fichier de primes incomplet ({report.row_count} lignes)")
    missing_cantons = sorted(OFFICIAL_CANTONS - report.cantons)
    if missing_cantons and path.stat().st_size > 1_000_000:
        report.reject("Cantons absents : " + ", ".join(missing_cantons))
    if report.quarantined_cantons:
        report.warnings.append(
            "Codes de canton non officiels écartés : "
            + ", ".join(f"{k} ({v})" for k, v in sorted(report.quarantined_cantons.items()))
        )
    if path.stat().st_size > 1_000_000:
        _check_region_shape(offers.values(), report)
    return report, offers


def _check_region_shape(offers: Iterable[Offer], report: ScanReport):
    seen: dict[str, set[str]] = {}
    for offer in offers:
        seen.setdefault(offer.canton, set()).add(offer.region)
    for canton in SINGLE_REGION_CANTONS:
        regions = seen.get(canton, set())
        if regions and regions != {"PR_REG_0"}:
            report.reject(f"{canton} devrait être en région unique PR_REG_0, trouvé {sorted(regions)}")
    for canton, numbers in MULTI_REGION_CANTONS.items():
        regions = seen.get(canton, set())
        expected = {f"PR_REG_{n}" for n in numbers}
        if regions and regions != expected:
            report.reject(f"Régions inattendues pour {canton} : {sorted(regions)}")


def load_region_xlsx(path: Path) -> dict[int, dict]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    regions = {}
    for row in sheet.iter_rows(values_only=True):
        if not row or not isinstance(row[1], int):
            continue
        bfs = int(row[1])
        canton = str(row[0]).strip().upper()
        number = int(row[3])
        if canton not in MULTI_REGION_CANTONS or number not in MULTI_REGION_CANTONS[canton]:
            raise ValueError(f"Région EDI inattendue : {canton} {bfs} {number}")
        regions[bfs] = {
            "canton": canton,
            "bfs": bfs,
            "commune": str(row[2]).strip(),
            "region_number": number,
            "region": f"PR_REG_{number}",
        }
    required = {5586: ("VD", 1), 351: ("BE", 1), 5401: ("VD", 2), 321: ("BE", 3)}
    for bfs, (canton, number) in required.items():
        item = regions.get(bfs)
        if not item or item["canton"] != canton or item["region_number"] != number:
            raise ValueError(f"Contrôle de région EDI échoué pour la commune OFS {bfs}")
    return regions


def load_localities_csv(text: str) -> list[dict]:
    rows = []
    reader = csv.DictReader(io.StringIO(text), delimiter=";")
    for row in reader:
        plz = re.sub(r"\D", "", row.get("PLZ4") or "")
        if len(plz) != 4:
            continue
        try:
            bfs = int(row["BFS-Nr"])
        except (TypeError, ValueError, KeyError):
            continue
        canton = (row.get("Kantonskürzel") or "").strip().upper()
        if canton not in OFFICIAL_CANTONS:
            continue
        rows.append({
            "locality": (row.get("Ortschaftsname") or "").strip(),
            "plz": plz,
            "commune": (row.get("Gemeindename") or "").strip(),
            "bfs": bfs,
            "canton": canton,
        })
    return rows


def load_localities_path(path: Path) -> list[dict]:
    if path.suffix.lower() == ".zip":
        with zipfile.ZipFile(path) as archive:
            name = next(item for item in archive.namelist() if item.endswith(".csv"))
            text = archive.read(name).decode("utf-8-sig")
        return load_localities_csv(text)
    return load_localities_csv(path.read_text(encoding="utf-8-sig"))


def load_tariff_names(path: Path) -> dict[tuple[str, str], str]:
    names = {}
    with _open_text(path) as handle:
        reader = csv.DictReader(handle)
        columns = _header_map(reader.fieldnames or [])
        if "name_fr" not in columns:
            return names
        for row in reader:
            if (row.get(columns.get("category", ""), "") or "").strip() != "MOD":
                continue
            insurer = str(row[columns["insurer"]]).strip()
            tarif = str(row[columns["tarif"]]).strip()
            names[(insurer, tarif)] = (row[columns["name_fr"]] or "").strip()
    return names


def load_catchment(path: Path) -> dict[tuple, Optional[set[int]]]:
    """None = toute la région. Un ensemble = communes OFS autorisées."""
    catchment: dict[tuple, Optional[set[int]]] = {}
    with _open_text(path) as handle:
        reader = csv.DictReader(handle)
        columns = _header_map(reader.fieldnames or [])
        for row in reader:
            canton = (row.get(columns.get("canton", ""), "") or "").strip().upper()
            if canton not in OFFICIAL_CANTONS:
                continue
            key = (
                str(row[columns["insurer"]]).strip(),
                canton,
                row[columns["region"]].strip(),
                row[columns["tarif"]].strip(),
            )
            restricted = (row.get(columns.get("restricted", ""), "") or "").strip().upper() == "Y"
            if not restricted:
                catchment[key] = None
                continue
            raw = row.get(columns.get("bfs_list", ""), "") or ""
            allowed = {int(part) for part in re.findall(r"\d+", raw)}
            current = catchment.get(key)
            if current is None and key in catchment:
                continue
            catchment[key] = (current or set()) | allowed
    return catchment


@dataclass
class Catalog:
    year: int
    offers: dict[tuple, Offer]
    localities: list[dict]
    regions_by_bfs: dict[int, dict]
    tariff_names: dict[tuple[str, str], str] = field(default_factory=dict)
    catchment: Optional[dict] = None
    source: dict = field(default_factory=dict)

    def offers_for(self, **filters) -> list[Offer]:
        found = []
        for offer in self.offers.values():
            if any(getattr(offer, key) != value for key, value in filters.items() if value is not None):
                continue
            found.append(self._named(offer))
        return found

    def _named(self, offer: Offer) -> Offer:
        french = self.tariff_names.get((offer.insurer_id, offer.tarif_id)) or offer.tarif_name_fr
        if french and french != offer.tarif_name_fr:
            offer.tarif_name_fr = french
        return offer


def resolve_residence(catalog: Catalog, *, plz: Optional[str] = None, commune: Optional[str] = None, canton: Optional[str] = None) -> dict:
    """Résout la région de primes sans inventer une commune."""
    plz_digits = re.sub(r"\D", "", plz or "")
    if len(plz_digits) == 5:
        return {"ok": False, "reason": "domicile_hors_suisse", "messages": ["Code postal français : personne exclue du comparateur AOS suisse."]}
    if len(plz_digits) not in {0, 4}:
        plz_digits = ""
    canton_code = (canton or "").strip().upper()
    if canton_code and canton_code not in OFFICIAL_CANTONS:
        canton_names = {
            "geneve": "GE", "genf": "GE", "vaud": "VD", "berne": "BE", "bern": "BE",
            "zurich": "ZH", "valais": "VS", "wallis": "VS", "fribourg": "FR", "freiburg": "FR",
            "neuchatel": "NE", "ticino": "TI", "tessin": "TI", "bale ville": "BS", "bale campagne": "BL",
            "lucerne": "LU", "luzern": "LU", "soleure": "SO", "argovie": "AG", "thurgovie": "TG",
            "saint gall": "SG", "st gallen": "SG", "grisons": "GR", "jura": "JU", "zoug": "ZG",
        }
        canton_code = canton_names.get(fold_name(canton), "")
    commune_key = fold_name(commune)
    candidates = []
    for row in catalog.localities:
        if plz_digits and row["plz"] != plz_digits:
            continue
        if canton_code and row["canton"] != canton_code:
            continue
        if commune_key:
            names = {fold_name(row["commune"]), fold_name(row["locality"])}
            if commune_key not in names:
                continue
        candidates.append(row)
    if not candidates and canton_code in SINGLE_REGION_CANTONS and not plz_digits and not commune_key:
        return _residence_ok(canton_code, "PR_REG_0", None, None, ambiguous=False, missing_commune=True)
    unique_bfs = {}
    for row in candidates:
        unique_bfs.setdefault(row["bfs"], row)
    if len(unique_bfs) == 1:
        row = next(iter(unique_bfs.values()))
        return _residence_from_row(catalog, row, ambiguous=False)
    if len(unique_bfs) > 1:
        resolved = []
        for row in unique_bfs.values():
            item = _residence_from_row(catalog, row, ambiguous=True)
            if item.get("ok"):
                resolved.append(item)
        regions = {item["region"] for item in resolved}
        cantons = {item["canton"] for item in resolved}
        if len(regions) == 1 and len(cantons) == 1 and resolved:
            merged = dict(resolved[0])
            merged["ambiguous_commune"] = True
            merged["commune"] = None
            merged["bfs"] = None
            merged["messages"] = ["Plusieurs communes correspondent. La région de primes est unique, la commune reste à préciser."]
            return merged
        return {
            "ok": False,
            "reason": "commune_ambigue",
            "messages": ["Plusieurs communes correspondent. Précisez la commune officielle avant de comparer."],
        }
    if commune_key:
        by_region = [
            item for item in catalog.regions_by_bfs.values()
            if fold_name(item["commune"]) == commune_key and (not canton_code or item["canton"] == canton_code)
        ]
        if len(by_region) == 1:
            item = by_region[0]
            return _residence_ok(item["canton"], item["region"], item["bfs"], item["commune"], ambiguous=False, missing_commune=False)
    if canton_code in SINGLE_REGION_CANTONS and (commune_key or plz_digits):
        return _residence_ok(canton_code, "PR_REG_0", None, commune or None, ambiguous=False, missing_commune=not commune_key)
    return {
        "ok": False,
        "reason": "residence_incomplete",
        "messages": ["Commune suisse ou code postal insuffisant pour déterminer la région de primes."],
    }


def _residence_from_row(catalog: Catalog, row: dict, *, ambiguous: bool) -> dict:
    canton = row["canton"]
    if canton in SINGLE_REGION_CANTONS:
        return _residence_ok(canton, "PR_REG_0", row["bfs"], row["commune"], ambiguous=ambiguous, missing_commune=False)
    region = catalog.regions_by_bfs.get(row["bfs"])
    if not region:
        return {
            "ok": False,
            "reason": "region_inconnue",
            "messages": [f"La commune OFS {row['bfs']} n'est pas dans l'annexe des régions de primes 2027."],
        }
    return _residence_ok(canton, region["region"], row["bfs"], row["commune"], ambiguous=ambiguous, missing_commune=False)


def _residence_ok(canton, region, bfs, commune, *, ambiguous: bool, missing_commune: bool) -> dict:
    return {
        "ok": True,
        "canton": canton,
        "region": region,
        "bfs": bfs,
        "commune": commune,
        "ambiguous_commune": ambiguous,
        "missing_commune": missing_commune,
        "messages": [],
    }


def _offer_allowed(catalog: Catalog, offer: Offer, bfs: Optional[int]) -> bool:
    if catalog.catchment is None:
        return True
    key = (offer.insurer_id, offer.canton, offer.region, offer.tarif_id)
    if key not in catalog.catchment:
        return False
    allowed = catalog.catchment[key]
    if allowed is None:
        return True
    return bfs is not None and bfs in allowed


def quote(catalog: Catalog, profile: dict) -> dict:
    """Compare les primes. N'affiche pas d'économie si la référence 2027 n'est pas unique."""
    missing = []
    year = int(profile.get("year") or catalog.year)
    residence = profile.get("residence") or {}
    if not residence.get("ok"):
        missing.append("commune")
    birth_year = profile.get("birth_year")
    if not birth_year:
        missing.append("annee_naissance")
        age_class = None
        subgroup = None
    else:
        age_class = age_class_for_birth_year(int(birth_year), year)
        subgroup = profile.get("age_subgroup") or ORDINARY_SUBGROUP[age_class]
    accident = profile.get("with_accident")
    if accident is None:
        missing.append("couverture_accidents")
    franchise = profile.get("franchise_chf")
    if franchise is None:
        missing.append("franchise")
    if missing or age_class is None or accident is None or franchise is None:
        return {
            "ok": False,
            "missing": missing,
            "age_class": age_class,
            "age_label": AGE_CLASS_FR.get(age_class or "", None),
            "offers": [],
            "savings_reliable": False,
            "messages": ["Comparaison bloquée : " + ", ".join(missing) + "."],
        }
    accident_value = accident_code(bool(accident))
    offers = [
        offer for offer in catalog.offers_for(
            canton=residence["canton"],
            region=residence["region"],
            age_class=age_class,
            age_subgroup=subgroup,
            accident=accident_value,
        )
        if _offer_allowed(catalog, offer, residence.get("bfs"))
    ]
    equivalent = [offer for offer in offers if offer.franchise_chf == int(franchise)]
    same_model = [offer for offer in equivalent if offer.tariftyp == profile.get("model_type")] if profile.get("model_type") else []
    reference = _find_reference(same_model, profile)
    savings_reliable = reference is not None and bool(profile.get("model_type")) and bool(profile.get("insurer_id"))
    reference_month = reference.premium_month if reference else None
    ranked = sorted(same_model if savings_reliable or profile.get("model_type") else equivalent, key=lambda item: (item.premium_month, item.insurer_id, item.tarif_id))
    comparison_a = ranked if profile.get("model_type") else []
    if not profile.get("model_type"):
        comparison_a = []
    alternatives = []
    for offer in sorted(offers, key=lambda item: (item.premium_month, item.insurer_id)):
        if offer.franchise_chf == int(franchise) and offer.tariftyp == profile.get("model_type"):
            continue
        alternatives.append(offer)
    def decorate(offer: Offer, *, equivalent_flag: bool) -> dict:
        payload = offer.public()
        payload["equivalent"] = equivalent_flag
        payload["participation_note"] = (
            f"Franchise {offer.franchise_chf} CHF. Plafond réglementaire de participation "
            f"(franchise + quote-part de 10 % plafonnée) : {payload['participation_ceiling_chf']} CHF par année. "
            "Ce plafond n'est pas une prévision de frais médicaux."
        )
        if savings_reliable and reference_month is not None:
            delta = money(reference_month - offer.premium_month)
            payload["savings_month"] = money_str(delta)
            payload["savings_year"] = money_str(annual_from_month(delta))
            payload["savings_reliable"] = equivalent_flag
        else:
            payload["savings_month"] = None
            payload["savings_year"] = None
            payload["savings_reliable"] = False
        return payload
    rows_a = [decorate(offer, equivalent_flag=True) for offer in comparison_a]
    rows_b = [decorate(offer, equivalent_flag=False) for offer in alternatives]
    top3 = rows_a[:3]
    best = next((row for row in rows_a if row["insurer_id"] != str(profile.get("insurer_id") or "") or row["tarif_id"] != profile.get("tarif_id")), rows_a[0] if rows_a else None)
    positive = None
    if savings_reliable and best and reference is not None:
        if not (best["insurer_id"] == reference.insurer_id and best["tarif_id"] == reference.tarif_id):
            delta = money(reference.premium_month - money(best["premium_month"]))
            if delta > 0:
                positive = {"month": money_str(delta), "year": money_str(annual_from_month(delta)), "offer": best}
    known_premium_year = profile.get("known_premium_year")
    known_note = None
    if profile.get("known_premium") is not None and known_premium_year != year:
        known_note = (
            f"La prime connue ({profile.get('known_premium')} CHF) n'est pas une prime {year} officielle "
            "et n'entre pas dans le calcul des économies."
        )
    return {
        "ok": True,
        "missing": [],
        "year": year,
        "canton": residence["canton"],
        "region": residence["region"],
        "commune": residence.get("commune"),
        "bfs": residence.get("bfs"),
        "birth_year": int(birth_year),
        "age_class": age_class,
        "age_label": AGE_CLASS_FR[age_class],
        "age_subgroup": subgroup,
        "with_accident": bool(accident),
        "franchise_chf": int(franchise),
        "model_type": profile.get("model_type"),
        "savings_reliable": bool(positive),
        "reference": reference.public() if reference else None,
        "reference_is_official_2027": reference is not None,
        "known_premium_ignored": known_note,
        "comparison_a": rows_a,
        "offers": rows_a or [
            decorate(offer, equivalent_flag=False)
            for offer in sorted(equivalent, key=lambda item: (item.premium_month, item.insurer_id, item.tarif_id))
        ],
        "comparison_b": rows_b,
        "top3": top3,
        "best_equivalent_alternative": positive,
        "medical_recommendation": False,
        "subsidies_included": False,
        "messages": [
            "La comparaison A conserve la franchise, la couverture accidents et le type de modèle.",
            "L'offre la moins chère n'est pas présentée comme la plus adaptée médicalement.",
            "Les subsides cantonaux individuels ne sont pas inclus.",
            known_note or "",
        ],
    }


def _find_reference(offers: list[Offer], profile: dict) -> Optional[Offer]:
    insurer = str(profile.get("insurer_id") or "").strip()
    model = profile.get("model_type")
    tarif_id = profile.get("tarif_id")
    if not insurer or not model:
        return None
    matches = [offer for offer in offers if offer.insurer_id == insurer and offer.tariftyp == model]
    if tarif_id:
        matches = [offer for offer in matches if offer.tarif_id == tarif_id]
    premiums = {offer.premium_month for offer in matches}
    if len(matches) == 1 or (matches and len(premiums) == 1 and tarif_id):
        return matches[0]
    if len(matches) == 1:
        return matches[0]
    return None


def birth_year_from_value(value) -> Optional[int]:
    if value is None or value == "":
        return None
    match = re.search(r"\b(19|20)\d{2}\b", str(value))
    if not match:
        return None
    return int(match.group(0))


def parse_accident(value) -> Optional[bool]:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return value
    text = fold_name(str(value))
    if text in {"1", "oui", "true", "avec", "avec accidents", "mit unf", "mit"}:
        return True
    if text in {"0", "non", "false", "sans", "sans accidents", "ohn unf", "ohne"}:
        return False
    return None


def parse_franchise_value(value) -> Optional[int]:
    if value is None or value == "":
        return None
    digits = re.sub(r"\D", "", str(value))
    if not digits:
        return None
    number = int(digits)
    allowed = {0, 100, 200, 300, 400, 500, 600, 1000, 1500, 2000, 2500}
    return number if number in allowed else None


def valid_email(value: Optional[str]) -> bool:
    return bool(value and _EMAIL_RE.match(value.strip()))


def render_email(client: dict, comparison: dict) -> dict:
    """Texte personnalisé. Aucune économie chiffrée si le calcul n'est pas fiable."""
    nom = (client.get("nom") or "").strip() or "—"
    sexe = fold_name(client.get("sexe") or client.get("civilite"))
    if sexe in {"femme", "madame", "mme"}:
        greeting = f"Bonjour Madame {nom},"
    elif sexe in {"homme", "monsieur", "m"}:
        greeting = f"Bonjour Monsieur {nom},"
    else:
        greeting = f"Bonjour Madame/Monsieur {nom},"
    subject = "Assurance-maladie 2027 – Réduisez potentiellement vos primes"
    current = comparison.get("reference")
    best = (comparison.get("best_equivalent_alternative") or {})
    if comparison.get("savings_reliable") and current and best.get("offer"):
        offer = best["offer"]
        body = "\n".join([
            greeting,
            "",
            "Les nouvelles primes d'assurance-maladie pour 2027 ont été publiées.",
            "",
            "Dans le cadre du suivi de votre dossier, nous avons effectué une comparaison des assurances-maladie disponibles dans votre commune de résidence.",
            "",
            "Selon les informations dont nous disposons :",
            "",
            f"Votre assurance actuelle : {current['insurer_name']} — {current['product_name']}",
            f"Votre prime mensuelle 2027 : {current['premium_month']} CHF",
            "",
            f"L'offre la plus économique identifiée : {offer['insurer_name']} — {offer['product_name']}",
            f"Nouvelle prime mensuelle : {offer['premium_month']} CHF",
            "",
            f"Économie mensuelle potentielle : {best['month']} CHF",
            f"Économie annuelle potentielle : {best['year']} CHF",
            "",
            "Cette comparaison concerne l'assurance obligatoire des soins et repose sur les informations disponibles à ce jour.",
            "Le choix d'un modèle d'assurance doit également tenir compte de ses conditions d'accès aux soins.",
            "",
            "Si vous souhaitez étudier cette possibilité et bénéficier de notre accompagnement dans vos démarches, nous vous invitons à répondre directement à cet e-mail.",
            "",
            "Nous restons à votre entière disposition.",
            "",
            "Avec nos meilleures salutations,",
            "",
            "L'équipe de l'Agence Mendes",
        ])
        kind = "economies"
    elif comparison.get("ok") and comparison.get("reference") and not comparison.get("savings_reliable"):
        subject = "Assurance-maladie 2027 – Comparaison de vos primes"
        body = "\n".join([
            greeting,
            "",
            "Les nouvelles primes d'assurance-maladie pour 2027 ont été publiées.",
            "",
            "Nous avons comparé les primes de l'assurance obligatoire des soins pour votre commune. Cette comparaison ne fait pas apparaître d'économie fiable à paramètres équivalents (même type de modèle, même franchise, même couverture accidents).",
            "",
            "Nous restons à votre disposition si vous souhaitez en discuter. Le choix d'un modèle doit aussi tenir compte de ses conditions d'accès aux soins.",
            "",
            "Avec nos meilleures salutations,",
            "",
            "L'équipe de l'Agence Mendes",
        ])
        kind = "aucune_economie"
    else:
        subject = "Assurance-maladie 2027 – Comparaison des primes de votre commune"
        body = "\n".join([
            greeting,
            "",
            "Les nouvelles primes d'assurance-maladie pour 2027 ont été publiées.",
            "",
            "Nous pouvons comparer les primes de l'assurance obligatoire des soins disponibles dans votre commune de résidence. Les informations actuellement au dossier ne permettent pas d'annoncer une économie personnalisée.",
            "",
            "Si vous souhaitez que nous préparions cette comparaison, vous pouvez répondre à cet e-mail.",
            "",
            "Avec nos meilleures salutations,",
            "",
            "L'équipe de l'Agence Mendes",
        ])
        kind = "sans_reference"
    if "diagnos" in body.casefold() or "maladie chronique" in body.casefold():
        raise RuntimeError("Le message ne doit pas contenir de donnée médicale")
    return {"subject": subject, "body": body, "kind": kind, "savings_claimed": kind == "economies"}


def real_send_enabled() -> bool:
    import os
    return os.environ.get("LAMAL_ALLOW_REAL_SEND", "").strip().lower() in {"1", "true", "yes"}


def commit_delivery(payload: dict, *, confirm: str, transport=None) -> dict:
    """N'appelle un transport que si l'envoi est explicitement autorisé et confirmé.

    Aucun transport SMTP n'est branché par défaut.
    """
    if not real_send_enabled():
        return {"ok": False, "status": "envoi_bloque", "detail": "L'envoi réel des e-mails LAMal est désactivé."}
    if confirm != "ENVOYER-LAMAL":
        return {"ok": False, "status": "validation_requise", "detail": "Validation explicite de l'administrateur requise."}
    if transport is None:
        return {"ok": False, "status": "envoi_bloque", "detail": "Aucun transport SMTP n'est branché pour ce module."}
    transport(payload)
    return {"ok": True, "status": "envoye"}
