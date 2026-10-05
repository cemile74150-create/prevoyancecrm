"""
Import identité clients Excel Suivi_3p → collection indépendante suivi_3p_clients.

Lie chaque fiche 3P à la base Clients hub via client_id (upsert sans source_import excel).
"""
from __future__ import annotations

import argparse
import asyncio
import io
import os
import sys
from datetime import datetime, timezone, date
from pathlib import Path
from typing import Any, Optional, Union

from dotenv import load_dotenv
from openpyxl import load_workbook

from suivi_3p import (
    COLLECTION_CLIENTS,
    DEFAULT_SUIVI_3P_STATUT,
    build_new_client,
)

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / ".env")

TENANT_USER_ID = os.environ.get("TENANT_USER_ID", "local-dev")

COL_NOM = 1
COL_PRENOM = 2
COL_NAISSANCE = 3
COL_ETAT_CIVIL = 5
COL_AGENT = 7
COL_NB_ENFANTS = 8
COL_CONJOINT_NOM = 10
COL_CONJOINT_PRENOM = 11
COL_CONJOINT_NAISSANCE = 12

ETAT_CIVIL_MAP = {
    "celibataire": "Célibataire",
    "célibataire": "Célibataire",
    "marie(e)": "Marié(e)",
    "marié(e)": "Marié(e)",
    "marie": "Marié(e)",
    "marié": "Marié(e)",
    "divorce(e)": "Divorcé(e)",
    "divorcé(e)": "Divorcé(e)",
    "veuf(ve)": "Veuf(ve)",
    "partenariat enregistre": "Partenariat enregistré",
    "partenariat enregistré": "Partenariat enregistré",
    "separe(e)": "Séparé(e)",
    "séparé(e)": "Séparé(e)",
}


def _clean_str(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _to_iso_date(v: Any) -> Optional[str]:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    if not s:
        return None
    if len(s) >= 10 and s[4] == "-" and s[7] == "-":
        return s[:10]
    return None


def _normalize_etat_civil(v: Any) -> Optional[str]:
    raw = _clean_str(v)
    if not raw:
        return None
    key = raw.casefold().replace("é", "e").replace("è", "e")
    for k, mapped in ETAT_CIVIL_MAP.items():
        if key == k.casefold().replace("é", "e").replace("è", "e") or raw.casefold() == k:
            return mapped
    for canonical in ETAT_CIVIL_MAP.values():
        if raw.casefold() == canonical.casefold():
            return canonical
    return raw


def _to_int(v: Any, default: int = 0) -> int:
    if v is None or v == "":
        return default
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return default


def _escape_regex(s: str) -> str:
    return "".join(f"\\{c}" if c in r".^$*+?{}[]\|()" else c for c in s)


def iter_suivi_rows(source: Union[str, Path, bytes, io.BytesIO]):
    if isinstance(source, (bytes, bytearray)):
        wb = load_workbook(io.BytesIO(source), data_only=True, read_only=True)
    elif isinstance(source, io.BytesIO):
        wb = load_workbook(source, data_only=True, read_only=True)
    else:
        wb = load_workbook(str(source), data_only=True, read_only=True)

    if "Suivi_3p" not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Feuille Suivi_3p introuvable. Feuilles: {wb.sheetnames}")

    ws = wb["Suivi_3p"]
    first = True
    for row in ws.iter_rows(min_row=1, max_col=15, values_only=True):
        if first:
            first = False
            continue
        nom = _clean_str(row[COL_NOM - 1] if len(row) >= COL_NOM else None)
        prenom = _clean_str(row[COL_PRENOM - 1] if len(row) >= COL_PRENOM else None)
        if not nom and not prenom:
            continue
        if not nom or not prenom:
            yield {"skip_reason": "nom_ou_prenom_manquant", "nom": nom, "prenom": prenom}
            continue
        yield {
            "nom": nom,
            "prenom": prenom,
            "date_naissance": _to_iso_date(row[COL_NAISSANCE - 1] if len(row) >= COL_NAISSANCE else None),
            "etat_civil": _normalize_etat_civil(row[COL_ETAT_CIVIL - 1] if len(row) >= COL_ETAT_CIVIL else None),
            "conseiller": _clean_str(row[COL_AGENT - 1] if len(row) >= COL_AGENT else None) or None,
            "nombre_enfants": _to_int(row[COL_NB_ENFANTS - 1] if len(row) >= COL_NB_ENFANTS else None),
            "conjoint_nom": _clean_str(row[COL_CONJOINT_NOM - 1] if len(row) >= COL_CONJOINT_NOM else None),
            "conjoint_prenom": _clean_str(row[COL_CONJOINT_PRENOM - 1] if len(row) >= COL_CONJOINT_PRENOM else None),
            "conjoint_date_naissance": _to_iso_date(
                row[COL_CONJOINT_NAISSANCE - 1] if len(row) >= COL_CONJOINT_NAISSANCE else None
            ),
        }
    wb.close()


async def import_suivi_3p(
    db,
    source: Union[str, Path, bytes, io.BytesIO],
    *,
    user_id: str = TENANT_USER_ID,
    dry_run: bool = False,
) -> dict:
    report = {
        "dry_run": dry_run,
        "created": 0,
        "skipped_existing": 0,
        "skipped_incomplete": 0,
        "errors": [],
        "details": [],
    }
    coll = db[COLLECTION_CLIENTS]
    rows = list(iter_suivi_rows(source))

    for row in rows:
        if row.get("skip_reason"):
            report["skipped_incomplete"] += 1
            continue
        try:
            q = {
                "user_id": user_id,
                "nom": {"$regex": f"^{_escape_regex(row['nom'])}$", "$options": "i"},
                "prenom": {"$regex": f"^{_escape_regex(row['prenom'])}$", "$options": "i"},
            }
            # date_naissance peut être chiffrée → matching en Python après decrypt
            from data_crypto_keys import encryption_configured
            from field_crypto import decrypt_client_fields, encrypt_client_fields

            candidates = await coll.find(q, {"_id": 0}).to_list(50)
            existing = None
            want_dob = row.get("date_naissance")
            for cand in candidates:
                plain = decrypt_client_fields(cand) if encryption_configured() else cand
                if want_dob and plain.get("date_naissance") != want_dob:
                    continue
                if not want_dob and len(candidates) > 1:
                    # Sans DOB et plusieurs homonymes : considérer comme existant (comportement prudent)
                    existing = cand
                    break
                existing = cand
                break
            if existing:
                report["skipped_existing"] += 1
                report["details"].append({"action": "skip_existing", "nom": row["nom"], "prenom": row["prenom"]})
                continue

            doc = build_new_client(
                user_id=user_id,
                prenom=row["prenom"],
                nom=row["nom"],
                date_naissance=row.get("date_naissance"),
                etat_civil=row.get("etat_civil"),
                nombre_enfants=row.get("nombre_enfants") or 0,
                conseiller=row.get("conseiller"),
                conjoint_prenom=row.get("conjoint_prenom"),
                conjoint_nom=row.get("conjoint_nom"),
                conjoint_date_naissance=row.get("conjoint_date_naissance"),
                statut=DEFAULT_SUIVI_3P_STATUT,
                source_import="suivi_3p_excel",
            )
            if not dry_run:
                from client_hub import upsert_hub_client_from_person

                hub, _ = await upsert_hub_client_from_person(
                    db,
                    {
                        "prenom": row["prenom"],
                        "nom": row["nom"],
                        "date_naissance": row.get("date_naissance"),
                        "etat_civil": row.get("etat_civil"),
                        "nombre_enfants": row.get("nombre_enfants") or 0,
                        "conjoint_prenom": row.get("conjoint_prenom"),
                        "conjoint_nom": row.get("conjoint_nom"),
                        "conjoint_date_naissance": row.get("conjoint_date_naissance"),
                    },
                    user_id=user_id,
                    module="suivi_3p",
                    conseiller=row.get("conseiller"),
                )
                if hub and hub.get("id"):
                    doc["client_id"] = hub["id"]
                to_store = encrypt_client_fields(doc) if encryption_configured() else doc
                await coll.insert_one(to_store)
            report["created"] += 1
            report["details"].append({
                "action": "created",
                "nom": row["nom"],
                "prenom": row["prenom"],
                "id": doc["id"],
            })
        except Exception as exc:
            report["errors"].append({"nom": row.get("nom"), "prenom": row.get("prenom"), "error": str(exc)})

    report["rows_read"] = len(rows)
    return report


async def _cli_main(path: Path, dry_run: bool):
    if os.environ.get("DATABASE_URL"):
        from postgres_mongo_compat import PostgresMongoCompatDB

        db = PostgresMongoCompatDB()
        report = await import_suivi_3p(db, path, dry_run=dry_run)
    else:
        from motor.motor_asyncio import AsyncIOMotorClient

        mongo_url = os.environ.get("MONGO_URL")
        db_name = os.environ.get("DB_NAME", "prevoyancecrm")
        if not mongo_url:
            print("MONGO_URL manquant", file=sys.stderr)
            sys.exit(1)
        client = AsyncIOMotorClient(mongo_url)
        db = client[db_name]
        report = await import_suivi_3p(db, path, dry_run=dry_run)
        client.close()
    print(f"rows_read={report['rows_read']}")
    print(f"created={report['created']} skipped_existing={report['skipped_existing']}")
    print(f"errors={len(report['errors'])}")


def main():
    parser = argparse.ArgumentParser(description="Import Suivi_3p → suivi_3p_clients")
    parser.add_argument(
        "excel",
        nargs="?",
        default=r"C:\Users\offic\Desktop\Suivi des 3p\Suivi_3P_Clients-Impots-V3.xlsm",
    )
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    path = Path(args.excel)
    if not path.exists():
        print(f"Fichier introuvable: {path}", file=sys.stderr)
        sys.exit(1)
    asyncio.run(_cli_main(path, args.dry_run))


if __name__ == "__main__":
    main()
