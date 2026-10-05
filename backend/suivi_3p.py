"""Module Suivi 3e Pilier — collection Mongo indépendante (suivi_3p_clients)."""
from __future__ import annotations

from datetime import datetime, timezone
from typing import Any, Optional

SUIVI_3P_STATUTS = [
    "À analyser",
    "Analyse en cours",
    "Offre envoyée",
    "Signé",
    "Refusé",
    "Sans suite",
]

DEFAULT_SUIVI_3P_STATUT = "À analyser"

# Suivi commercial — statut du rendez-vous (différent du statut dossier)
RDV_STATUTS = ["En cours", "Signé", "Annulé"]
DEFAULT_RDV_STATUT = None


def normalize_rdv_statut(value) -> Optional[str]:
    """Normalise le statut RDV ; None / vide = non renseigné."""
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    for allowed in RDV_STATUTS:
        if s.casefold() == allowed.casefold():
            return allowed
    return None


def rdv_compte_rendu_filled(value) -> bool:
    return bool(str(value or "").strip())


def is_rdv_suivi_complete(doc: dict) -> bool:
    """Suivi RDV complété = statut ET compte rendu renseignés."""
    if not isinstance(doc, dict):
        return False
    return bool(normalize_rdv_statut(doc.get("rdv_statut"))) and rdv_compte_rendu_filled(
        doc.get("rdv_compte_rendu")
    )


def needs_rdv_suivi_relance(doc: dict, *, today=None) -> bool:
    """
    True si une relance e-mail J+1 doit être envoyée :
    - RDV pris avec date
    - date_rdv + 1 jour <= today
    - suivi incomplet (statut OU compte rendu manquant)
    - relance pas déjà envoyée pour ce RDV
    """
    from datetime import date as date_cls, timedelta

    if not isinstance(doc, dict):
        return False
    if not doc.get("rdv_pris"):
        return False
    raw = doc.get("date_rdv")
    if not raw:
        return False
    day = str(raw)[:10]
    try:
        y, m, d = int(day[0:4]), int(day[5:7]), int(day[8:10])
        rdv_day = date_cls(y, m, d)
    except (TypeError, ValueError, IndexError):
        return False
    ref = today or date_cls.today()
    if ref < (rdv_day + timedelta(days=1)):
        return False
    if is_rdv_suivi_complete(doc):
        return False
    if doc.get("rdv_suivi_relance_sent_at"):
        return False
    return True

ANALYSE_DOC_CATEGORY = "Analyse 3e Pilier"
COURRIER_KIND = "courrier_optimisation_fiscale"
COURRIER_DISPLAY_PREFIX = "Courrier optimisation fiscale"
GAIN_OPTIM_MIN_EXCLUSIVE = 250.0
COLLECTION_CLIENTS = "suivi_3p_clients"
COLLECTION_DOCS = "suivi_3p_documents"
COLLECTION_PENDING = "suivi_3p_pending_documents"


def is_analyse_pdf_doc(doc: dict) -> bool:
    """True pour un PDF d'analyse 3P (classique ou Fortune), hors courrier Word."""
    if not isinstance(doc, dict) or doc.get("is_deleted"):
        return False
    if (doc.get("kind") or "").strip() == COURRIER_KIND:
        return False
    if (doc.get("category") or "").strip() == ANALYSE_DOC_CATEGORY:
        return True
    label = (doc.get("display_label") or "").casefold()
    folder = (doc.get("source_folder") or "").casefold()
    if "optimisation fiscale" in label or "fortune" in folder:
        return True
    return False


def is_offre_envoyee_eligible(client: dict, docs: Optional[list] = None) -> bool:
    """Offre envoyée = au moins un PDF d'analyse + gain fiscal > 250 CHF."""
    docs = docs or []
    has_pdf = any(is_analyse_pdf_doc(d) for d in docs)
    if not has_pdf:
        return False
    gain = parse_gain(client.get("gain_fiscal_estime"))
    if gain is None:
        # repli éventuel sur gain extrait d'un PDF
        for d in docs:
            if not is_analyse_pdf_doc(d):
                continue
            g = parse_gain(d.get("extracted_gain_fiscal"))
            if g is not None and float(g) > GAIN_OPTIM_MIN_EXCLUSIVE:
                return True
        return False
    return float(gain) > GAIN_OPTIM_MIN_EXCLUSIVE


def normalize_person_name(value: Any) -> str:
    """Normalise nom/prénom pour matching (accents, espaces, casse)."""
    import unicodedata

    s = str(value or "").strip()
    if not s:
        return ""
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return " ".join(s.casefold().split())


def person_match_key(nom: Any, prenom: Any) -> str:
    return f"{normalize_person_name(nom)}|{normalize_person_name(prenom)}"


def parse_analyse_pdf_filename(filename: str) -> tuple[Optional[str], Optional[str]]:
    """
    Parse noms de PDF d'analyse :
      - 'NOM_Prenom.pdf' / 'DOS SANTOS_Ivan.pdf' / 'GOUCHAULT_Jean Paul.pdf'
      - 'NOM Prenom.pdf' / 'DOS SANTOS GONCALVES Leonel.pdf' / 'DA SILVA José.pdf'
        (dernier mot = prénom, le reste = nom)
    Retourne (nom, prenom) ou (None, None).
    """
    from pathlib import Path

    stem = Path(filename or "").stem.strip()
    if not stem:
        return None, None
    # ignore placeholders
    if normalize_person_name(stem) in {"nom|prenom", "nom prenom"}:
        return None, None

    if "_" in stem:
        nom_part, prenom_part = stem.split("_", 1)
        nom = nom_part.strip()
        prenom = prenom_part.strip()
    else:
        parts = stem.split()
        if len(parts) < 2:
            return None, None
        # Convention dossier « Analyse Fortune » : NOM… Prénom (dernier token)
        prenom = parts[-1].strip()
        nom = " ".join(parts[:-1]).strip()

    if not nom or not prenom:
        return None, None
    if normalize_person_name(nom) == "nom" and normalize_person_name(prenom) == "prenom":
        return None, None
    return nom, prenom


def build_client_name_index(clients: list) -> dict:
    """
    Map match_key → list[client].
    Inclut aussi les clés inversées prénom|nom pour tolérance.
    """
    index: dict = {}
    for c in clients:
        keys = {
            person_match_key(c.get("nom"), c.get("prenom")),
            person_match_key(c.get("prenom"), c.get("nom")),
        }
        # conjoint éventuel (PDF parfois au nom du conjoint)
        cp = c.get("conjoint_prenom")
        cn = c.get("conjoint_nom")
        if cp and cn:
            keys.add(person_match_key(cn, cp))
            keys.add(person_match_key(cp, cn))
        for k in keys:
            if not k or k == "|":
                continue
            index.setdefault(k, []).append(c)
    return index


def find_clients_for_pdf_name(filename: str, index: dict) -> tuple[Optional[str], Optional[str], list]:
    """Retourne (parsed_nom, parsed_prenom, matches)."""
    nom, prenom = parse_analyse_pdf_filename(filename)
    if not nom or not prenom:
        return nom, prenom, []
    key = person_match_key(nom, prenom)
    matches = index.get(key) or []
    # dédupliquer par id
    seen = set()
    unique = []
    for c in matches:
        cid = c.get("id")
        if cid in seen:
            continue
        seen.add(cid)
        unique.append(c)
    return nom, prenom, unique


def normalize_statut(statut: Optional[str]) -> str:
    s = (statut or DEFAULT_SUIVI_3P_STATUT).strip()
    return s if s in SUIVI_3P_STATUTS else DEFAULT_SUIVI_3P_STATUT


def parse_gain(value: Any) -> Optional[float]:
    """Parse un montant CHF (1 839, 1'839, 1839.5, '1 839 CHF', etc.)."""
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    import re

    s = str(value).strip()
    s = re.sub(r"(?i)\bchf\b", "", s).strip()
    s = (
        s.replace("'", "")
        .replace("\u2019", "")
        .replace("\u00a0", "")
        .replace(" ", "")
        .replace(",", ".")
    )
    try:
        return float(s)
    except (TypeError, ValueError):
        return None


def extract_economie_fiscale_from_pdf(pdf_bytes: bytes) -> Optional[float]:
    """
    Lit le bandeau « ÉCONOMIE FISCALE ESTIMÉE » (bas du PDF d'analyse Mendes)
    et retourne le montant en CHF.
    """
    import re
    import unicodedata

    if not pdf_bytes:
        return None

    text = ""
    try:
        import fitz  # pymupdf

        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        text = "\n".join((page.get_text("text") or "") for page in doc)
    except Exception:
        try:
            import io
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(pdf_bytes))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        except Exception:
            return None

    if not text.strip():
        return None

    norm = unicodedata.normalize("NFKD", text)
    norm = "".join(ch for ch in norm if not unicodedata.combining(ch))

    patterns = [
        # Bandeau principal en bas du PDF
        r"ECONOMIE\s+FISCALE\s+ESTIMEE\s*[:\-]?\s*([0-9][0-9\s'’.,]*)\s*(?:CHF)?",
        # Variante « Economie estimée » plus haut dans le document
        r"Economie\s+estimee\s*[:\-]?\s*([0-9][0-9\s'’.,]*)\s*(?:CHF)?",
    ]
    for pat in patterns:
        m = re.search(pat, norm, re.IGNORECASE)
        if not m:
            continue
        gain = parse_gain(m.group(1))
        if gain is not None:
            return gain
    return None


def suivi_3p_base_query(user, conseiller: Optional[str] = None) -> dict:
    """Filtre tenant + scoping conseiller (même logique que la prévoyance)."""
    from access_control import (
        TENANT_USER_ID,
        is_global_viewer,
        conseiller_scope_mongo_filter,
    )
    import re

    q: dict = {"user_id": TENANT_USER_ID}
    if not is_global_viewer(user):
        q.update(conseiller_scope_mongo_filter(user))
    elif conseiller and conseiller != "all":
        if conseiller == "Non attribué":
            q["$and"] = [{
                "$or": [
                    {"conseiller": {"$exists": False}},
                    {"conseiller": None},
                    {"conseiller": ""},
                    {"conseiller": "Non attribué"},
                ]
            }]
        else:
            q["conseiller"] = {"$regex": f"^{re.escape(conseiller.strip())}$", "$options": "i"}
    return q


def can_access_suivi_client(user, client: dict) -> bool:
    from access_control import can_access_client

    return can_access_client(user, client)

def is_analyse_fortune_doc(doc: dict) -> bool:
    """True pour un PDF Analyse Fortune (dossier Fortune), pas une analyse 3e pilier classique."""
    if not isinstance(doc, dict) or doc.get("is_deleted"):
        return False
    kind = (doc.get("kind") or "").strip()
    if kind == COURRIER_KIND:
        return False
    # Libellé / type explicite
    analyse_kind = str(doc.get("analyse_kind") or "").strip().casefold()
    if analyse_kind in {"fortune", "analyse_fortune"}:
        return True
    if analyse_kind in {"3p", "3e_pilier", "pilier", "frontalier"}:
        return False
    folder = (doc.get("source_folder") or "").casefold()
    name = (doc.get("original_filename") or "").casefold()
    label = (doc.get("display_label") or "").casefold()
    # Dossier / nom « Fortune » uniquement — ne pas confondre avec
    # « Analyse d'optimisation fiscale » (PDF 3e pilier classiques / frontaliers).
    if "fortune" in folder or "fortune" in name:
        return True
    if label.startswith("analyse fortune") or label == "analyse fortune":
        return True
    return False


def is_analyse_3p_classic_doc(doc: dict) -> bool:
    """True pour une Analyse 3e Pilier classique (hors Fortune et hors courrier)."""
    if not isinstance(doc, dict) or doc.get("is_deleted"):
        return False
    kind = (doc.get("kind") or "").strip()
    if kind == COURRIER_KIND:
        return False
    if is_analyse_fortune_doc(doc):
        return False
    analyse_kind = str(doc.get("analyse_kind") or "").strip().casefold()
    if analyse_kind in {"3p", "3e_pilier", "pilier", "frontalier"}:
        return True
    if (doc.get("category") or "").strip() == ANALYSE_DOC_CATEGORY:
        return True
    folder = (doc.get("source_folder") or "").casefold()
    if "frontalier" in folder or "3e pilier" in folder or "3p" in folder:
        return True
    label = (doc.get("display_label") or "").casefold()
    if "3e pilier" in label or "optimisation fiscale" in label:
        return True
    return False


def analyse_summary_from_docs(docs: Optional[list] = None) -> dict:
    """
    Résumé pour la liste Suivi 3P :
      has_analyse_3p, has_analyse_fortune, analyse_fortune_doc_id, analyse_fortune_filename
    """
    has_3p = False
    fortune_doc = None
    for d in docs or []:
        if not isinstance(d, dict) or d.get("is_deleted"):
            continue
        if is_analyse_fortune_doc(d):
            if fortune_doc is None or (d.get("created_at") or "") >= (fortune_doc.get("created_at") or ""):
                fortune_doc = d
        elif is_analyse_3p_classic_doc(d):
            has_3p = True
    return {
        "has_analyse_3p": has_3p,
        "has_analyse_fortune": fortune_doc is not None,
        "analyse_fortune_doc_id": fortune_doc.get("id") if fortune_doc else None,
        "analyse_fortune_filename": (
            fortune_doc.get("original_filename")
            or fortune_doc.get("display_label")
            or "Analyse Fortune.pdf"
        )
        if fortune_doc
        else None,
    }


def classify_residence_suivi(client: dict) -> Optional[str]:
    """
    Statut géographique Suivi 3P : « Frontalier » | « Suisse » | None.

    Priorité à l'adresse / pays (infer_country) — pas uniquement le NPA.
    Le champ stocké ``frontalier`` sert de repli si le pays n'est pas déductible.
    """
    from courrier_word import infer_country

    country = infer_country(client or {})
    if country == "France":
        return "Frontalier"
    if country == "Suisse":
        return "Suisse"

    flag = str((client or {}).get("frontalier") or "").strip().casefold()
    if flag in {"oui", "yes", "true", "1", "frontalier"}:
        return "Frontalier"
    if flag in {"suisse", "switzerland", "ch"}:
        return "Suisse"
    return None


def serialize_suivi_client(doc: dict) -> dict:
    from data_crypto_keys import EncryptionNotConfigured
    from field_crypto import decrypt_client_fields

    try:
        doc = decrypt_client_fields(doc)
    except EncryptionNotConfigured:
        raise
    statut = normalize_statut(doc.get("statut"))
    gain = parse_gain(doc.get("gain_fiscal_estime"))
    date_a = doc.get("date_derniere_analyse")
    if date_a:
        date_a = str(date_a)[:10]
    else:
        date_a = None
    conjoint = doc.get("conjoint")
    if not conjoint:
        cp = (doc.get("conjoint_prenom") or "").strip()
        cn = (doc.get("conjoint_nom") or "").strip()
        conjoint = f"{cp} {cn}".strip() or None
    residence = classify_residence_suivi(doc)
    return {
        "id": doc.get("id"),
        "prenom": doc.get("prenom") or "",
        "nom": doc.get("nom") or "",
        "date_naissance": doc.get("date_naissance"),
        "etat_civil": doc.get("etat_civil"),
        "nombre_enfants": doc.get("nombre_enfants") or 0,
        "conseiller": doc.get("conseiller"),
        "conjoint": conjoint,
        "conjoint_prenom": doc.get("conjoint_prenom"),
        "conjoint_nom": doc.get("conjoint_nom"),
        "conjoint_date_naissance": doc.get("conjoint_date_naissance"),
        "email": doc.get("email"),
        "telephone": doc.get("telephone"),
        "adresse": doc.get("adresse"),
        "adresse_complement": doc.get("adresse_complement"),
        "npa": doc.get("npa"),
        "ville": doc.get("ville"),
        "pays": doc.get("pays") or doc.get("pays_residence"),
        "frontalier": doc.get("frontalier"),
        "residence_suivi": residence,
        "sexe": doc.get("sexe") or doc.get("civilite"),
        "civilite": doc.get("civilite") or doc.get("sexe"),
        "conjoint_sexe": doc.get("conjoint_sexe") or doc.get("sexe_conjoint"),
        "statut": statut,
        "statut_suivi": statut,
        "gain_fiscal_estime": gain,
        "date_derniere_analyse": date_a,
        "client_contacte": bool(doc.get("client_contacte")),
        "rdv_pris": bool(doc.get("rdv_pris")),
        "date_rdv": (str(doc.get("date_rdv"))[:10] if doc.get("date_rdv") else None),
        "rdv_statut": normalize_rdv_statut(doc.get("rdv_statut")),
        "rdv_compte_rendu": (str(doc.get("rdv_compte_rendu")).strip() if doc.get("rdv_compte_rendu") else None) or None,
        "rdv_suivi_complete": is_rdv_suivi_complete(doc),
        "rdv_suivi_relance_sent_at": doc.get("rdv_suivi_relance_sent_at"),
        "notes": doc.get("notes"),
        "created_at": doc.get("created_at"),
        "updated_at": doc.get("updated_at"),
        "source_import": doc.get("source_import"),
        "client_id": doc.get("client_id"),
    }


def build_new_client(
    *,
    user_id: str,
    prenom: str,
    nom: str,
    date_naissance: Optional[str] = None,
    etat_civil: Optional[str] = None,
    nombre_enfants: int = 0,
    conseiller: Optional[str] = None,
    conjoint: Optional[str] = None,
    conjoint_prenom: Optional[str] = None,
    conjoint_nom: Optional[str] = None,
    conjoint_date_naissance: Optional[str] = None,
    email: Optional[str] = None,
    telephone: Optional[str] = None,
    adresse: Optional[str] = None,
    npa: Optional[str] = None,
    ville: Optional[str] = None,
    adresse_complement: Optional[str] = None,
    pays: Optional[str] = None,
    sexe: Optional[str] = None,
    statut: str = DEFAULT_SUIVI_3P_STATUT,
    gain_fiscal_estime: Optional[float] = None,
    date_derniere_analyse: Optional[str] = None,
    notes: Optional[str] = None,
    source_import: Optional[str] = None,
    client_id: Optional[str] = None,
) -> dict:
    import uuid

    now = datetime.now(timezone.utc).isoformat()
    cp = (conjoint_prenom or "").strip()
    cn = (conjoint_nom or "").strip()
    conj = (conjoint or "").strip() or (f"{cp} {cn}".strip() or None)
    return {
        "id": client_id or str(uuid.uuid4()),
        "user_id": user_id,
        "prenom": (prenom or "").strip(),
        "nom": (nom or "").strip(),
        "date_naissance": (str(date_naissance)[:10] if date_naissance else None),
        "etat_civil": etat_civil,
        "nombre_enfants": int(nombre_enfants or 0),
        "conseiller": (conseiller or "").strip() or None,
        "conjoint": conj,
        "conjoint_prenom": cp or None,
        "conjoint_nom": cn or None,
        "conjoint_date_naissance": (str(conjoint_date_naissance)[:10] if conjoint_date_naissance else None),
        "email": email,
        "telephone": telephone,
        "adresse": (adresse or "").strip() or None,
        "adresse_complement": (adresse_complement or "").strip() or None,
        "npa": (npa or "").strip() or None,
        "ville": (ville or "").strip() or None,
        "pays": (pays or "").strip() or None,
        "sexe": (sexe or "").strip() or None,
        "statut": normalize_statut(statut),
        "gain_fiscal_estime": parse_gain(gain_fiscal_estime),
        "date_derniere_analyse": (str(date_derniere_analyse)[:10] if date_derniere_analyse else None),
        "client_contacte": False,
        "rdv_pris": False,
        "date_rdv": None,
        "rdv_statut": None,
        "rdv_compte_rendu": None,
        "rdv_suivi_relance_sent_at": None,
        "notes": (str(notes).strip() if notes else None) or None,
        "source_import": source_import,
        "created_at": now,
        "updated_at": now,
    }


def compute_stats(rows: list) -> dict:
    from datetime import date as date_cls

    by_statut: dict = {s: 0 for s in SUIVI_3P_STATUTS}
    gain_total = 0.0
    analyses_realisees = 0
    signes = 0
    rdv_pris = 0
    clients_contactes = 0
    upcoming: list = []
    today = date_cls.today().isoformat()

    for r in rows:
        st = r["statut_suivi"]
        by_statut[st] = by_statut.get(st, 0) + 1
        if st == "Signé":
            signes += 1
        if st != "À analyser" or r.get("date_derniere_analyse"):
            analyses_realisees += 1
        g = r.get("gain_fiscal_estime")
        if g is not None and st not in ("Refusé", "Sans suite"):
            try:
                gain_total += float(g)
            except (TypeError, ValueError):
                pass
        if r.get("client_contacte"):
            clients_contactes += 1
        if r.get("rdv_pris"):
            rdv_pris += 1
            d = r.get("date_rdv")
            if d:
                day = str(d)[:10]
                if day >= today:
                    upcoming.append({
                        "id": r.get("id"),
                        "prenom": r.get("prenom") or "",
                        "nom": r.get("nom") or "",
                        "conseiller": r.get("conseiller") or "—",
                        "date_rdv": day,
                    })

    upcoming.sort(key=lambda x: (x["date_rdv"], (x.get("nom") or "").casefold(), (x.get("prenom") or "").casefold()))

    frontaliers = sum(1 for r in rows if r.get("residence_suivi") == "Frontalier")
    suisses = sum(1 for r in rows if r.get("residence_suivi") == "Suisse")

    return {
        "total": len(rows),
        "analyses_realisees": analyses_realisees,
        "gain_fiscal_total": round(gain_total, 2),
        "contrats_signes": signes,
        "rdv_pris": rdv_pris,
        "clients_contactes": clients_contactes,
        "a_contacter": max(0, len(rows) - clients_contactes),
        "offres_envoyees": by_statut.get("Offre envoyée", 0),
        "upcoming_rdvs": upcoming[:50],
        "by_statut": by_statut,
        "frontaliers": frontaliers,
        "suisses": suisses,
    }


def filter_suivi_3p_rows(
    rows: list,
    *,
    statut: Optional[str] = None,
    filtre: Optional[str] = None,
    geo: Optional[str] = None,
    q: Optional[str] = None,
) -> list:
    """Filtres liste Suivi 3P (statut, KPI, géo Frontaliers/Suisses, recherche texte)."""
    out = list(rows)

    if statut and statut != "all":
        out = [r for r in out if r.get("statut_suivi") == statut]

    filtre_norm = (filtre or "").strip().lower()
    if filtre_norm == "analyses":
        out = [
            r for r in out
            if r.get("statut_suivi") != "À analyser" or r.get("date_derniere_analyse")
        ]
    elif filtre_norm == "gain":
        out = [
            r for r in out
            if r.get("gain_fiscal_estime") is not None
            and float(r.get("gain_fiscal_estime") or 0) != 0
        ]
    elif filtre_norm in ("signes", "contrats"):
        out = [r for r in out if r.get("statut_suivi") == "Signé"]
    elif filtre_norm in ("rdv", "rdv_pris"):
        out = [r for r in out if r.get("rdv_pris")]
    elif filtre_norm in ("contactes", "contact"):
        out = [r for r in out if r.get("client_contacte")]
    elif filtre_norm in ("a_contacter", "acontacter"):
        out = [r for r in out if not r.get("client_contacte")]
    elif filtre_norm in ("frontaliers", "frontalier"):
        out = [r for r in out if r.get("residence_suivi") == "Frontalier"]
    elif filtre_norm in ("suisses", "suisse"):
        out = [r for r in out if r.get("residence_suivi") == "Suisse"]

    # Filtre géo dédié (combinable avec les KPI via `filtre`)
    geo_norm = (geo or "").strip().lower()
    if geo_norm in ("frontaliers", "frontalier"):
        out = [r for r in out if r.get("residence_suivi") == "Frontalier"]
    elif geo_norm in ("suisses", "suisse"):
        out = [r for r in out if r.get("residence_suivi") == "Suisse"]

    if q and q.strip():
        ql = q.strip().casefold()
        out = [
            r for r in out
            if ql in (r.get("prenom") or "").casefold()
            or ql in (r.get("nom") or "").casefold()
            or ql in f"{r.get('prenom', '')} {r.get('nom', '')}".casefold()
            or ql in f"{r.get('nom', '')} {r.get('prenom', '')}".casefold()
            or ql in (r.get("conseiller") or "").casefold()
            or ql in (r.get("conjoint") or "").casefold()
            or ql in (r.get("conjoint_prenom") or "").casefold()
            or ql in (r.get("conjoint_nom") or "").casefold()
            or ql in (r.get("residence_suivi") or "").casefold()
        ]

    out.sort(key=lambda r: ((r.get("nom") or "").casefold(), (r.get("prenom") or "").casefold()))
    return out


def _courrier_doc_from_list(docs: Optional[list]) -> Optional[dict]:
    best = None
    for d in docs or []:
        if not isinstance(d, dict) or d.get("is_deleted"):
            continue
        if (d.get("kind") or "").strip() != COURRIER_KIND:
            continue
        if best is None or (d.get("created_at") or "") >= (best.get("created_at") or ""):
            best = d
    return best


def courrier_sent_date(docs: Optional[list]) -> Optional[str]:
    """Date d'envoi du courrier (ISO YYYY-MM-DD) depuis le document courrier."""
    doc = _courrier_doc_from_list(docs)
    if not doc:
        return None
    for key in ("generated_at", "created_at", "updated_at"):
        raw = doc.get(key)
        if raw:
            return str(raw)[:10]
    return None


def has_courrier_sent(docs: Optional[list]) -> bool:
    return _courrier_doc_from_list(docs) is not None


def needs_phone_followup(row: dict, docs: Optional[list] = None) -> bool:
    """Client à contacter par la téléphoniste (courrier envoyé, pas encore joint)."""
    if row.get("client_contacte"):
        return False
    statut = row.get("statut_suivi") or row.get("statut") or ""
    if statut == "Offre envoyée":
        return True
    if has_courrier_sent(docs):
        return True
    return False


def phone_followup_statut_label(row: dict, docs: Optional[list] = None) -> str:
    """Libellé statut pour l'export téléphoniste."""
    if needs_phone_followup(row, docs):
        return "À appeler"
    if row.get("rdv_pris"):
        return "RDV pris"
    if row.get("client_contacte"):
        return "Contacté"
    return row.get("statut_suivi") or row.get("statut") or ""


def format_date_fr(iso: Optional[str]) -> str:
    if not iso:
        return ""
    s = str(iso).strip()[:10]
    parts = s.split("-")
    if len(parts) == 3:
        return f"{parts[2]}.{parts[1]}.{parts[0]}"
    return s


def format_gain_fiscal_export(value: Any) -> str:
    """Montant enregistré tel quel (pas de recalcul), formaté pour Excel."""
    g = parse_gain(value)
    if g is None:
        return ""
    # Format suisse : 2450 → 2'450
    as_int = int(round(g))
    raw = f"{as_int:,}".replace(",", "'")
    return f"{raw} CHF"


def _client_display_name(row: dict) -> str:
    return f"{(row.get('prenom') or '').strip()} {(row.get('nom') or '').strip()}".strip()


def _conjoint_display_name(row: dict) -> str:
    cp = (row.get("conjoint_prenom") or "").strip()
    cn = (row.get("conjoint_nom") or "").strip()
    if cp or cn:
        return f"{cp} {cn}".strip()
    return (row.get("conjoint") or "").strip()


def _excel_sheet_name(name: str, used: set) -> str:
    invalid = set(":\\/?*[]")
    base = "".join(c for c in (name or "Non attribué").strip() if c not in invalid) or "Non attribué"
    base = base[:31]
    candidate = base
    n = 1
    while candidate in used:
        suffix = f" ({n})"
        candidate = base[: max(1, 31 - len(suffix))] + suffix
        n += 1
    used.add(candidate)
    return candidate


EXPORT_COLUMN_DEFS = [
    ("conseiller", "Conseiller / Agent"),
    ("client", "Nom et prénom du client"),
    ("gain_fiscal", "Gain fiscal proposé (CHF)"),
    ("conjoint", "Nom et prénom du conjoint"),
    ("telephone", "Téléphone"),
    ("email", "E-mail"),
    ("date_courrier", "Date d'envoi du courrier"),
    ("statut", "Statut du suivi"),
    ("date_suivi", "Date du prochain suivi / échéance"),
    ("notes", "Commentaire / note"),
]
EXPORT_COLUMNS = [label for _, label in EXPORT_COLUMN_DEFS]
DEFAULT_EXPORT_COLUMN_KEYS = [key for key, _ in EXPORT_COLUMN_DEFS]

EXPORT_SCOPES = frozenset({"a_appeler", "liste_filtree", "selection"})
EXPORT_GROUPINGS = frozenset({"par_conseiller", "feuille_unique"})


def normalize_export_columns(columns: Optional[list] = None) -> list:
    allowed = {k for k, _ in EXPORT_COLUMN_DEFS}
    if not columns:
        return list(DEFAULT_EXPORT_COLUMN_KEYS)
    out = []
    for c in columns:
        key = (str(c or "")).strip()
        if key in allowed and key not in out:
            out.append(key)
    return out or list(DEFAULT_EXPORT_COLUMN_KEYS)


def row_in_export_scope(row: dict, docs: Optional[list], scope: str) -> bool:
    scope = (scope or "a_appeler").strip().lower()
    if scope == "liste_filtree" or scope == "selection":
        return True
    return needs_phone_followup(row, docs)


def build_suivi_3p_export_xlsx(
    rows: list,
    docs_map: Optional[dict] = None,
    *,
    scope: str = "a_appeler",
    grouping: str = "par_conseiller",
    columns: Optional[list] = None,
) -> bytes:
    """Génère un classeur Excel Suivi 3P (périmètre, regroupement et colonnes configurables)."""
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    scope = (scope or "a_appeler").strip().lower()
    if scope not in EXPORT_SCOPES:
        scope = "a_appeler"
    grouping = (grouping or "par_conseiller").strip().lower()
    if grouping not in EXPORT_GROUPINGS:
        grouping = "par_conseiller"
    col_keys = normalize_export_columns(columns)
    col_labels = [label for key, label in EXPORT_COLUMN_DEFS if key in col_keys]
    statut_col_idx = col_keys.index("statut") + 1 if "statut" in col_keys else None

    docs_map = docs_map or {}
    export_rows = []
    for row in rows:
        cid = row.get("id")
        docs = docs_map.get(cid) or []
        if not row_in_export_scope(row, docs, scope):
            continue
        conseiller = (row.get("conseiller") or "").strip() or "Non attribué"
        export_rows.append({
            "conseiller": conseiller,
            "client": _client_display_name(row),
            "gain_fiscal": format_gain_fiscal_export(row.get("gain_fiscal_estime")),
            "conjoint": _conjoint_display_name(row),
            "telephone": row.get("telephone") or "",
            "email": row.get("email") or "",
            "date_courrier": format_date_fr(courrier_sent_date(docs)),
            "statut": phone_followup_statut_label(row, docs),
            "date_suivi": format_date_fr(row.get("date_rdv")),
            "notes": (row.get("notes") or "").strip(),
            "nom_sort": (row.get("nom") or "").casefold(),
            "prenom_sort": (row.get("prenom") or "").casefold(),
            "client_id": cid or "",
        })

    export_rows.sort(
        key=lambda r: (
            r["conseiller"].casefold(),
            r["nom_sort"],
            r["prenom_sort"],
        )
    )

    wb = Workbook()
    wb.remove(wb.active)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="002FA7")
    call_fill = PatternFill("solid", fgColor="FEF3C7")
    used_names: set = set()

    def _write_sheet(ws, sheet_rows: list) -> None:
        ws.append(col_labels)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for r in sheet_rows:
            ws.append([r.get(k, "") for k in col_keys])
            if statut_col_idx and r.get("statut") == "À appeler":
                row_idx = ws.max_row
                for col in range(1, len(col_keys) + 1):
                    ws.cell(row=row_idx, column=col).fill = call_fill
        ws.freeze_panes = "A2"
        for col_idx in range(1, len(col_keys) + 1):
            letter = ws.cell(row=1, column=col_idx).column_letter
            max_len = max(
                len(str(ws.cell(row=row_idx, column=col_idx).value or ""))
                for row_idx in range(1, ws.max_row + 1)
            )
            ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 48)

    if not export_rows:
        ws = wb.create_sheet("Aucun client")
        _write_sheet(ws, [])
    elif grouping == "feuille_unique":
        ws = wb.create_sheet("Export")
        _write_sheet(ws, export_rows)
    else:
        by_conseiller: dict = {}
        for r in export_rows:
            by_conseiller.setdefault(r["conseiller"], []).append(r)
        for conseiller in sorted(by_conseiller.keys(), key=lambda x: x.casefold()):
            ws = wb.create_sheet(_excel_sheet_name(conseiller, used_names))
            _write_sheet(ws, by_conseiller[conseiller])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_suivi_3p_phone_export_xlsx(
    rows: list,
    docs_map: Optional[dict] = None,
) -> bytes:
    """Rétrocompatibilité : export téléphoniste par défaut."""
    return build_suivi_3p_export_xlsx(rows, docs_map, scope="a_appeler", grouping="par_conseiller")
