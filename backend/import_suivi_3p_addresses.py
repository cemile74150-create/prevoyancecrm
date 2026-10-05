"""
Import Excel d'adresses → mise à jour de `suivi_3p_clients`.
Ne crée pas de nouveaux clients (pas de doublons).
"""
from __future__ import annotations

import io
import re
from datetime import datetime, timezone
from typing import Any, Dict, List, Tuple

from openpyxl import load_workbook

from suivi_3p import COLLECTION_CLIENTS, person_match_key


HEADER_ALIASES = {
    "id": ("id", "client_id", "clientid", "uuid"),
    "nom": ("nom", "name", "lastname", "last_name", "family_name", "name"),
    "prenom": ("prenom", "prénom", "firstname", "first_name", "vorname"),
    "adresse": ("adresse", "address", "rue", "street", "strasse", "domicile"),
    "npa": ("npa", "cp", "code postal", "zip", "plz"),
    "ville": ("ville", "city", "localite", "localité", "ort", "lieu"),
}


def _norm_header(v: Any) -> str:
    s = str(v or "").strip().lower()
    for a, b in (("é", "e"), ("è", "e"), ("ê", "e"), ("à", "a"), ("ô", "o"), ("û", "u")):
        s = s.replace(a, b)
    return re.sub(r"\s+", " ", s)


def _clean(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def _map_headers(row: tuple) -> Dict[str, int]:
    mapping: Dict[str, int] = {}
    for idx, cell in enumerate(row):
        h = _norm_header(cell)
        if not h:
            continue
        for field, aliases in HEADER_ALIASES.items():
            if field in mapping:
                continue
            if h in aliases or any(a == h or (len(a) > 3 and a in h) for a in aliases):
                mapping[field] = idx
                break
    return mapping


def _cell(row: tuple, mapping: Dict[str, int], key: str) -> str:
    idx = mapping.get(key)
    if idx is None or idx >= len(row):
        return ""
    return _clean(row[idx])


def split_full_address(raw: str) -> Tuple[str, str, str]:
    """
    Découpe une adresse complète type :
      'Rue de Lausanne, 63 - 1202 Genève'
      '20, Place des Aviateurs - F 74580 VIRY'
      'Rue Philippe Plantamour 33 1201 Genève'
    → (rue, npa, ville)
    """
    text = re.sub(r"\s+", " ", (raw or "").strip())
    if not text:
        return "", "", ""

    # «... - F 01170 Crozet» ou «... - 1202 Genève»
    m = re.search(r"^(.*?)\s*[-–]\s*(?:F\s*)?(\d{4,5})\s+(.+)$", text, re.IGNORECASE)
    if m:
        return m.group(1).strip(" ,"), m.group(2), m.group(3).strip()

    # «... 1201 Genève» (NPA CH 4 chiffres en fin)
    m = re.search(r"^(.*?)\s+(\d{4})\s+([A-Za-zÀ-ÿ][\wÀ-ÿ'’.\- ]+)$", text)
    if m:
        return m.group(1).strip(" ,"), m.group(2), m.group(3).strip()

    return text, "", ""


def parse_address_rows(file_bytes: bytes) -> Tuple[List[Dict[str, str]], List[str]]:
    errors: List[str] = []
    try:
        wb = load_workbook(filename=io.BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as e:
        return [], [f"Fichier Excel illisible : {e}"]

    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        return [], ["Fichier vide"]

    mapping = _map_headers(header or ())
    if "adresse" not in mapping and "npa" not in mapping and "ville" not in mapping:
        return [], ["Colonnes adresse introuvables (attendu : Adresse / NPA / Ville)"]
    if "nom" not in mapping and "prenom" not in mapping and "id" not in mapping:
        return [], ["Colonnes d'identification introuvables (Nom/Prénom ou Id)"]

    out: List[Dict[str, str]] = []
    for i, row in enumerate(rows_iter, start=2):
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        item = {
            "id": _cell(row, mapping, "id"),
            "nom": _cell(row, mapping, "nom"),
            "prenom": _cell(row, mapping, "prenom"),
            "adresse": _cell(row, mapping, "adresse"),
            "npa": _cell(row, mapping, "npa"),
            "ville": _cell(row, mapping, "ville"),
            "row_num": str(i),
        }
        # Adresse.xlsx : une seule colonne « Adresse » complète → découper
        if item["adresse"] and not item["npa"] and not item["ville"]:
            rue, npa, ville = split_full_address(item["adresse"])
            item["adresse"] = rue or item["adresse"]
            item["npa"] = npa
            item["ville"] = ville
        if not any([item["adresse"], item["npa"], item["ville"]]):
            errors.append(f"Ligne {i} : aucune donnée d'adresse")
            continue
        if not any([item["id"], item["nom"], item["prenom"]]):
            errors.append(f"Ligne {i} : identification manquante (nom/prénom ou id)")
            continue
        out.append(item)
    return out, errors


async def apply_suivi_3p_address_import(db, user_id: str, file_bytes: bytes) -> Dict[str, Any]:
    """
    Met à jour adresse/npa/ville des clients Suivi 3P existants.
    Ne crée jamais de nouveau client.
    """
    rows, parse_errors = parse_address_rows(file_bytes)
    updated: List[Dict[str, str]] = []
    errors: List[str] = list(parse_errors)
    skipped = 0

    clients = await db[COLLECTION_CLIENTS].find({"user_id": user_id}, {"_id": 0}).to_list(20000)
    by_id = {c.get("id"): c for c in clients if c.get("id")}
    by_name: Dict[str, list] = {}
    for c in clients:
        key = person_match_key(c.get("nom"), c.get("prenom"))
        if key and key != "|":
            by_name.setdefault(key, []).append(c)
        # tolérance prénom|nom inversé
        key2 = person_match_key(c.get("prenom"), c.get("nom"))
        if key2 and key2 != "|" and key2 != key:
            by_name.setdefault(key2, []).append(c)

    for row in rows:
        match = None
        if row["id"] and row["id"] in by_id:
            match = by_id[row["id"]]
        else:
            key = person_match_key(row["nom"], row["prenom"])
            cands = by_name.get(key) or []
            # dédupliquer
            seen = set()
            unique = []
            for c in cands:
                cid = c.get("id")
                if cid in seen:
                    continue
                seen.add(cid)
                unique.append(c)
            if len(unique) == 1:
                match = unique[0]
            elif len(unique) > 1:
                errors.append(
                    f"Ligne {row['row_num']} : plusieurs clients pour {row['prenom']} {row['nom']}"
                )
                continue
            else:
                errors.append(
                    f"Ligne {row['row_num']} : client introuvable ({row['prenom']} {row['nom'] or row['id']})".strip()
                )
                continue

        if match is None:
            errors.append(f"Ligne {row['row_num']} : client introuvable")
            continue

        updates = {}
        if row["adresse"]:
            updates["adresse"] = row["adresse"]
        if row["npa"]:
            updates["npa"] = row["npa"]
        if row["ville"]:
            updates["ville"] = row["ville"]
        if not updates:
            skipped += 1
            continue

        from data_crypto_keys import encryption_configured
        from field_crypto import decrypt_client_fields, patch_encrypt_sensitive

        match_plain = decrypt_client_fields(match) if encryption_configured() else match
        same = all(str(match_plain.get(k) or "") == str(v) for k, v in updates.items())
        if same:
            skipped += 1
            continue

        plain_snapshot = dict(updates)
        to_set = patch_encrypt_sensitive(updates) if encryption_configured() else dict(updates)
        to_set["updated_at"] = datetime.now(timezone.utc).isoformat()
        await db[COLLECTION_CLIENTS].update_one(
            {"id": match["id"], "user_id": user_id},
            {"$set": to_set},
        )
        updated.append({
            "client_id": match["id"],
            "nom": match.get("nom") or "",
            "prenom": match.get("prenom") or "",
            "adresse": plain_snapshot.get("adresse", match_plain.get("adresse") or ""),
            "npa": plain_snapshot.get("npa", match_plain.get("npa") or ""),
            "ville": plain_snapshot.get("ville", match_plain.get("ville") or ""),
        })

    return {
        "updated_count": len(updated),
        "updated": updated[:200],
        "error_count": len(errors),
        "errors": errors[:200],
        "skipped": skipped,
        "rows_read": len(rows),
    }
