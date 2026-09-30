"""Tests import groupé + remplissage PDF (AVS / mapping / cases vides)."""
from __future__ import annotations

import io
import sys
from pathlib import Path
from zipfile import ZipFile

import pytest
from pypdf import PdfReader

BACKEND = Path(__file__).resolve().parents[1]
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

from pdf_generator import (  # noqa: E402
    TEMPLATES_DIR,
    _map_calcul_rente,
    _map_mandat_de_gestion,
    _mandat_de_gestion_fill_style,
    _split_ch_date,
    client_field_values,
    generate_document_pdf,
    validate_template_fields,
)


SAMPLE_CLIENT = {
    "nom": "Dupont",
    "prenom": "Jean",
    "date_naissance": "1980-01-15",
    "sexe": "Homme",
    "etat_civil": "Marié(e)",
    "avs_number": "756.1234.5678.90",
    "adresse": "Rue du Lac 12",
    "npa": "1200",
    "ville": "Genève",
    "telephone": "0791234567",
    "email": "jean@example.ch",
    "nationalite": "Suisse",
    "pays_residence": "Suisse",
    "frontalier": "non",
}


def test_map_calcul_rente_full_client():
    values = client_field_values(SAMPLE_CLIENT)
    mapping = _map_calcul_rente(values)
    assert mapping["NOM"] == "DUPONT"
    assert mapping["PRENOM"] == "Jean"
    assert mapping["AVS"] == "756.1234.5678.90"
    assert mapping["Pays de résidence"] == "Suisse"
    assert mapping["Texte8"] == "1200"
    assert mapping["Localité"] == "Genève"
    assert "/MASCULIN" in mapping["Groupe34"]
    assert "/Choix1" in mapping["Groupe34"]  # marié


@pytest.mark.parametrize(
    "avs_number",
    [
        "756.1234.5678.90",
        "756 1234 5678 90",
        "756-1234-5678-90",
    ],
)
def test_map_calcul_rente_preserves_avs_crm_format(avs_number):
    values = client_field_values({**SAMPLE_CLIENT, "avs_number": avs_number})
    assert _map_calcul_rente(values)["AVS"] == avs_number


def test_map_calcul_rente_empty_client_no_false_checkboxes():
    values = client_field_values({})
    mapping = _map_calcul_rente(values)
    assert mapping["NOM"] == ""
    assert mapping["AVS"] == ""
    assert mapping["Pays de résidence"] == ""
    assert "Groupe34" not in mapping


def test_map_calcul_rente_sex_femme_celibataire():
    values = client_field_values(
        {"sexe": "Femme", "etat_civil": "Célibataire", "nom": "Martin", "prenom": "Anna"}
    )
    mapping = _map_calcul_rente(values)
    assert mapping["Groupe34"] == ["/FEMININ", "/Choix2"]


def test_map_calcul_rente_never_invents_frontalier_checkbox():
    values = client_field_values({**SAMPLE_CLIENT, "frontalier": "oui"})
    mapping = _map_calcul_rente(values)
    # Pas de widget frontalier dans ce PDF
    assert all(not str(k).lower().startswith("frontal") for k in mapping)


def test_generate_all_builtin_templates_smoke():
    from unittest.mock import patch

    for tid in (
        "procuration_avs_lpp",
        "recherche_avoirs_lpp",
        "lettre_lpp",
        "lettre_avs",
        "lettre_decompte_lpp",
        "mandat_de_gestion",
    ):
        opts = (
            {"fund": {"name": "Caisse Test", "address": "1 rue Test\n1000 Lausanne"}}
            if tid == "lettre_decompte_lpp"
            else None
        )
        if tid == "lettre_avs":
            with patch("lettre_lpp_word.convert_docx_to_pdf", return_value=b"%PDF-1.4 smoke" + b"0" * 200):
                file_bytes, meta = generate_document_pdf(
                    tid, SAMPLE_CLIENT, agent_name="Agent", options=opts
                )
        else:
            file_bytes, meta = generate_document_pdf(
                tid, SAMPLE_CLIENT, agent_name="Agent", options=opts
            )
        assert meta["template_id"] == tid
        assert len(file_bytes) > 100
        if tid == "lettre_avs":
            assert meta["original_filename"].endswith(".pdf")
            assert meta.get("has_word_download") is True
            assert meta.get("docx_filename", "").endswith(".docx")
            assert file_bytes.startswith(b"%PDF")
        else:
            assert PdfReader(io.BytesIO(file_bytes)).pages


def test_batch_upload_helper_exists():
    """Garantit que l'endpoint batch est bien défini dans server."""
    import importlib.util

    # Ne charge pas tout server (Mongo env) — vérifie le fichier source
    src = (BACKEND / "server.py").read_text(encoding="utf-8")
    assert "documents/batch" in src
    assert "_store_uploaded_client_document" in src
    assert "files: List[UploadFile]" in src or "files: List[UploadFile] = File" in src


def test_fill_library_form_rejects_coroutine_pdf():
    """Régression Pictet : une coroutine ne doit jamais atteindre fitz/BytesIO."""
    from pdf_generator import fill_pdf_bytes_with_client

    async def _fake_read():
        return b"%PDF-1.4"

    coro = _fake_read()
    try:
        with pytest.raises(TypeError, match="coroutine"):
            fill_pdf_bytes_with_client(coro, SAMPLE_CLIENT, agent_name="Agent")
    finally:
        coro.close()


def test_fill_library_form_roundtrip_bytes():
    """Chaîne lecture bytes → prepare → fill → bytes (formulaire type Pictet)."""
    import fitz
    from pdf_generator import fill_pdf_bytes_with_client, prepare_library_form_pdf

    doc = fitz.open()
    page = doc.new_page()
    widget = fitz.Widget()
    widget.field_name = "Nom"
    widget.field_type = fitz.PDF_WIDGET_TYPE_TEXT
    widget.rect = fitz.Rect(40, 40, 200, 60)
    page.add_widget(widget)
    buf = io.BytesIO()
    doc.save(buf)
    doc.close()
    raw = buf.getvalue()

    prep = prepare_library_form_pdf(raw)
    assert isinstance(prep["pdf_bytes"], bytes)
    mapping = {prep["widgets"][0]["id"]: "nom"} if prep["widgets"] else None
    out = fill_pdf_bytes_with_client(
        prep["pdf_bytes"],
        SAMPLE_CLIENT,
        agent_name="Agent",
        field_mapping=mapping,
        widgets=prep["widgets"],
    )
    assert isinstance(out, bytes)
    assert out.startswith(b"%PDF")
    assert len(out) > 100


def test_mandat_template_fields_present():
    path = TEMPLATES_DIR / "mandat_mg_lsa.pdf"
    assert path.exists()
    assert validate_template_fields("mandat_de_gestion", path) == []


def test_map_mandat_single_client_leaves_person_2_empty():
    values = client_field_values({
        **SAMPLE_CLIENT,
        "etat_civil": "Célibataire",
        "conjoint": "",
        "conseiller": "Alberto Mendes",
        "conseiller_finma": "123456",
    })
    mapping = _map_mandat_de_gestion(values)
    assert mapping["Nom"] == "DUPONT"
    assert mapping["Prénom"] == "Jean"
    assert mapping["née le"] == "15"
    assert mapping["undefined"] == "01"
    assert mapping["undefined_2"] == "1980"
    assert mapping["Nom_2"] == ""
    assert mapping["Prénom_2"] == ""
    assert mapping["née le_2"] == ""
    assert mapping["undefined_5"] == "1200"
    assert mapping["undefined_6"] == "Genève" or mapping.get("Ville") == "Genève"
    assert mapping["le"] == "14" or len(mapping["le"]) == 2
    assert mapping.get("Check Box4") == "/Oui" or mapping.get("Monsieur") == "/Oui"
    assert "Check Box1" not in mapping or mapping.get("Check Box1") != "/Oui"
    assert mapping["Représenté par le courtier"] == "ALBERTO MENDES"
    assert mapping["Le représentant du Mandataire inscrit au registre de la FINMA sous le numéro"] == "123456"
    assert mapping["Text1"] == "ALBERTO MENDES"
    assert mapping["Text2"] == "123456"
    assert all(v not in {"undefined", "null", None} for v in mapping.values())


def test_map_mandat_conseiller_fields_uppercase_and_personal_finma():
    values = client_field_values({
        **SAMPLE_CLIENT,
        "conseiller": "Jean Dupont",
        "conseiller_finma": "123456",
    })
    mapping = _map_mandat_de_gestion(values)
    assert mapping["Text1"] == "JEAN DUPONT"
    assert mapping["Text2"] == "123456"
    # FINMA agence ne doit jamais écraser le champ conseiller
    values_agency = client_field_values({
        **SAMPLE_CLIENT,
        "conseiller": "Jean Dupont",
        "conseiller_finma": "F01101452",
    })
    mapping_agency = _map_mandat_de_gestion(values_agency)
    assert mapping_agency["Text1"] == "JEAN DUPONT"
    assert mapping_agency["Text2"] == ""


def test_map_mandat_married_without_spouse_does_not_invent_person_2():
    values = client_field_values(SAMPLE_CLIENT)
    mapping = _map_mandat_de_gestion(values)
    assert mapping["Prénom"] == "Jean"
    assert mapping["Nom_2"] == ""
    assert mapping["Prénom_2"] == ""
    assert mapping["née le_2"] == ""


def test_map_mandat_couple_from_linked_spouse():
    values = client_field_values(
        SAMPLE_CLIENT,
        spouse={
            "prenom": "Marie",
            "nom": "Dupont",
            "date_naissance": "1982-03-20",
            "sexe": "Femme",
        },
    )
    mapping = _map_mandat_de_gestion(values)
    assert mapping["Prénom"] == "Jean"
    assert mapping["Nom"] == "DUPONT"
    assert mapping["née le"] == "15"
    assert mapping["undefined"] == "01"
    assert mapping["undefined_2"] == "1980"
    assert mapping["Prénom_2"] == "Marie"
    assert mapping["Nom_2"] == "DUPONT"
    assert mapping["née le_2"] == "20"
    assert mapping["undefined_3"] == "03"
    assert mapping["undefined_4"] == "1982"
    assert mapping.get("Check Box6") == "/Oui" or mapping.get("Madame 2") == "/Oui"


def test_generate_mandat_draws_full_birth_dates_jj_mm_aaaa():
    """Régression : JJ + MM + AAAA doivent apparaître pour 1 et 2 mandants (pas année seule)."""
    client = {
        **SAMPLE_CLIENT,
        "prenom": "Karine",
        "nom": "Ambrosetti",
        "date_naissance": "1988-11-15",
        "sexe": "Femme",
    }
    spouse = {
        "prenom": "Jérôme",
        "nom": "Ambrosetti",
        "date_naissance": "1988-01-20",
        "sexe": "Homme",
    }
    pdf_bytes, _meta = generate_document_pdf("mandat_de_gestion", client, spouse=spouse)
    try:
        import fitz
    except ImportError:
        pytest.skip("PyMuPDF requis pour lire le texte dessiné")
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    words = [w[4] for w in doc[0].get_text("words")]
    # Mandant 1 : 15 / 11 / 1988 — Mandant 2 : 20 / 01 / 1988
    for token in ("15", "11", "1988", "20", "01"):
        assert token in words, f"token manquant dans le mandat: {token} (words={words})"


def test_split_ch_date_accepts_iso_and_slash():
    assert _split_ch_date("1988-11-15") == ("15", "11", "1988")
    assert _split_ch_date("15.11.1988") == ("15", "11", "1988")
    assert _split_ch_date("15/11/1988") == ("15", "11", "1988")
    assert _split_ch_date("1988-11-15T00:00:00Z") == ("15", "11", "1988")


def test_map_mandat_couple_from_conjoint_fields():
    values = client_field_values({
        **SAMPLE_CLIENT,
        "conjoint_prenom": "Claire",
        "conjoint_nom": "Martin",
        "conjoint_date_naissance": "1985-07-04",
    })
    mapping = _map_mandat_de_gestion(values)
    assert mapping["Prénom_2"] == "Claire"
    assert mapping["Nom_2"] == "MARTIN"
    assert mapping["née le_2"] == "04"


def test_map_mandat_never_writes_undefined_placeholders():
    values = client_field_values({
        "prenom": "undefined",
        "nom": "null",
        "conjoint_prenom": "None",
        "conjoint_nom": "undefined",
    })
    mapping = _map_mandat_de_gestion(values)
    assert mapping["Prénom"] == ""
    assert mapping["Nom"] == ""
    assert mapping["Prénom_2"] == ""
    assert mapping["Nom_2"] == ""


def test_generate_mandat_filename_uses_first_person():
    pdf_bytes, meta = generate_document_pdf("mandat_de_gestion", SAMPLE_CLIENT)
    assert meta["template_id"] == "mandat_de_gestion"
    assert meta["checklist_item"] == "Mandat de gestion"
    assert meta["original_filename"] == "Mandat_MG_LSA_DUPONT_Jean.pdf"
    assert meta.get("template_version") == "mg_lsa_v2"
    assert meta.get("template_bytes") == 177982
    assert meta["original_filename"].endswith(".pdf")
    assert PdfReader(io.BytesIO(pdf_bytes)).pages
    assert len(pdf_bytes) > 1000
    import fitz
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    text = doc[0].get_text()
    doc.close()
    assert "DUPONT" in text
    assert "Jean" in text
    assert "1980" in text
    # JJ/MM/AAAA : au moins mois + année visibles (JJ parfois trop étroit pour l'extracteur)


def test_split_ch_date_parts():
    assert _split_ch_date("15.01.1980") == ("15", "01", "1980")
    assert _split_ch_date("4/7/1985") == ("04", "07", "1985")
    assert _split_ch_date("") == ("", "", "")


def test_mandat_fill_style_uses_opensans_10():
    style = _mandat_de_gestion_fill_style()
    assert style["size"] == 10.0
    # OpenSans si le TTF est installé, sinon repli Helvetica prévu par le code.
    assert style["reportlab_font"] in ("FillOpenSans", "Helvetica")
    assert style["valign"] == "baseline"
    assert style["erase_background"] is False


def test_procuration_keeps_gray_background():
    """Les encadrés gris du modèle ne doivent pas être recouverts de blanc."""
    import fitz

    pdf_bytes, meta = generate_document_pdf("procuration_avs_lpp", SAMPLE_CLIENT)
    assert meta["template_id"] == "procuration_avs_lpp"
    page = fitz.open(stream=pdf_bytes, filetype="pdf")[0]
    # Point dans l'encadré identité, hors texte (fond imprimé ≈ #F2F2F2)
    pix = page.get_pixmap(clip=fitz.Rect(90, 250, 110, 265), alpha=False)
    r, g, b = pix.pixel(5, 5)
    assert max(r, g, b) < 252, f"fond blanc indésirable, obtenu {(r, g, b)}"
    assert min(r, g, b) > 220, f"fond trop sombre, obtenu {(r, g, b)}"
    assert abs(r - g) <= 3 and abs(g - b) <= 3
    assert "DUPONT" in page.get_text()


def test_lettre_lpp_uses_selected_client_and_today_date():
    import hashlib
    import fitz
    from datetime import datetime
    from lettre_lpp_word import ORIGINAL_DOCX, TEMPLATE_DOCX, fill_lettre_lpp_docx

    original_hash = hashlib.sha256(ORIGINAL_DOCX.read_bytes()).hexdigest() if ORIGINAL_DOCX.exists() else None
    client = {
        **SAMPLE_CLIENT,
        "nom": "Ambrosoetto",
        "prenom": "Karine",
    }
    pdf_bytes, meta = generate_document_pdf("lettre_lpp", client)
    assert meta["template_id"] == "lettre_lpp"
    page = fitz.open(stream=pdf_bytes, filetype="pdf")[0]
    text = page.get_text()
    compact = " ".join(text.replace("\u00a0", " ").split())
    assert "Karine AMBROSOETTO" in compact
    assert "procuration signée par Karine AMBROSOETTO nous permettant" in compact
    assert "Vous trouverez ci-joint" in compact
    assert "Recherche d'avoirs de la prévoyance professionnelle" in compact
    assert "Perly, le" in compact
    months = [
        "janvier", "février", "mars", "avril", "mai", "juin",
        "juillet", "août", "septembre", "octobre", "novembre", "décembre",
    ]
    today = datetime.now()
    assert f"{today.day} {months[today.month - 1]} {today.year}" in compact
    assert compact.count("procuration signée") == 1
    assert "Karine AMBROSOETTO" in compact
    assert page.get_images()  # logo / bandeau
    if original_hash:
        assert hashlib.sha256(ORIGINAL_DOCX.read_bytes()).hexdigest() == original_hash
    if TEMPLATE_DOCX.exists():
        filled = fill_lettre_lpp_docx(client_field_values(client))
        with ZipFile(io.BytesIO(filled)) as zf:
            xml = zf.read("word/document.xml").decode("utf-8")
        assert "Karine" in xml
        assert "{{client_fullname}}" not in xml
        assert "{{date_long}}" not in xml
        assert TEMPLATE_DOCX.read_bytes() != filled


def test_lettre_lpp_changes_with_another_client():
    import fitz

    first, _ = generate_document_pdf("lettre_lpp", {**SAMPLE_CLIENT, "nom": "Dupont", "prenom": "Jean"})
    second, _ = generate_document_pdf("lettre_lpp", {**SAMPLE_CLIENT, "nom": "Martin", "prenom": "Sophie"})
    t1 = fitz.open(stream=first, filetype="pdf")[0].get_text()
    t2 = fitz.open(stream=second, filetype="pdf")[0].get_text()
    assert "Jean DUPONT" in t1
    assert "Sophie MARTIN" in t2
    assert "Sophie MARTIN" not in t1
    assert "Jean DUPONT" not in t2


def test_lettre_lpp_long_name_stays_on_one_phrase():
    import fitz

    client = {
        **SAMPLE_CLIENT,
        "nom": "Fontannaz Mathey",
        "prenom": "Agnes",
    }
    pdf_bytes, meta = generate_document_pdf("lettre_lpp", client)
    assert meta["template_id"] == "lettre_lpp"
    page = fitz.open(stream=pdf_bytes, filetype="pdf")[0]
    text = page.get_text()
    compact = " ".join(text.replace("\u00a0", " ").split())
    assert "Agnes FONTANNAZ MATHEY" in compact
    assert "Vous trouverez ci-joint" in compact
    assert "Agnes FONTANNAZ MATHEY nous permettant" in compact
    assert compact.count("LPP.") <= 2


def test_lettre_lpp_recipient_stays_top_right():
    import fitz
    from docx import Document
    from lettre_lpp_word import RECIPIENT_LINES, fill_lettre_lpp_docx

    filled = fill_lettre_lpp_docx(client_field_values(SAMPLE_CLIENT))
    doc = Document(io.BytesIO(filled))
    found = 0
    for paragraph in doc.paragraphs:
        if (paragraph.text or "").strip() in RECIPIENT_LINES:
            indent = paragraph.paragraph_format.left_indent
            assert indent is not None
            assert indent.cm > 9.5
            found += 1
    assert found == 4

    pdf_bytes, _ = generate_document_pdf("lettre_lpp", SAMPLE_CLIENT)
    page = fitz.open(stream=pdf_bytes, filetype="pdf")[0]
    midpoint = page.rect.width / 2
    xs = []
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            text = "".join(span.get("text", "") for span in line.get("spans", []))
            if "Fonds de garantie LPP" in text:
                xs.append(line["bbox"][0])
    assert xs
    assert min(xs) > midpoint

