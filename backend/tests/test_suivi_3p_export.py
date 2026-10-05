from io import BytesIO

from openpyxl import load_workbook

from suivi_3p import (
    build_suivi_3p_export_xlsx,
    build_suivi_3p_phone_export_xlsx,
    filter_suivi_3p_rows,
    needs_phone_followup,
    normalize_export_columns,
    phone_followup_statut_label,
)


def _row(**kwargs):
    base = {
        "id": "1",
        "prenom": "Jean",
        "nom": "Dupont",
        "conseiller": "MENDES Alberto",
        "statut_suivi": "Offre envoyée",
        "client_contacte": False,
        "rdv_pris": False,
        "date_rdv": None,
        "email": "jean@example.com",
        "telephone": "+41 79 000 00 00",
        "notes": "Rappeler en matinée",
    }
    base.update(kwargs)
    return base


def test_needs_phone_followup_offre_envoyee_non_contacte():
    row = _row()
    assert needs_phone_followup(row, []) is True
    assert phone_followup_statut_label(row, []) == "À appeler"


def test_needs_phone_followup_excludes_contacte():
    row = _row(client_contacte=True)
    assert needs_phone_followup(row, []) is False
    assert phone_followup_statut_label(row, []) == "Contacté"


def test_filter_respects_statut_and_search():
    rows = [
        _row(id="1", nom="Dupont"),
        _row(id="2", nom="Martin", statut_suivi="Signé"),
    ]
    out = filter_suivi_3p_rows(rows, statut="Offre envoyée", q="Martin")
    assert len(out) == 0
    out = filter_suivi_3p_rows(rows, statut="Offre envoyée")
    assert len(out) == 1
    assert out[0]["nom"] == "Dupont"


def test_export_xlsx_one_sheet_per_conseiller():
    rows = [
        _row(id="1", conseiller="Agent A", gain_fiscal_estime=2450),
        _row(id="2", prenom="Marie", nom="Curie", conseiller="Agent B", gain_fiscal_estime=1850),
        _row(id="3", client_contacte=True, conseiller="Agent A"),
    ]
    docs_map = {
        "1": [{"kind": "courrier_optimisation_fiscale", "created_at": "2026-01-15T10:00:00"}],
        "2": [{"kind": "courrier_optimisation_fiscale", "created_at": "2026-02-01T10:00:00"}],
    }
    data = build_suivi_3p_phone_export_xlsx(rows, docs_map)
    wb = load_workbook(BytesIO(data))
    assert set(wb.sheetnames) == {"Agent A", "Agent B"}
    ws_a = wb["Agent A"]
    assert ws_a.max_row == 2
    # colonnes: conseiller, client, gain_fiscal, conjoint, telephone, email, date_courrier, statut…
    assert ws_a.cell(1, 3).value == "Gain fiscal proposé (CHF)"
    assert ws_a.cell(2, 3).value == "2'450 CHF"
    assert ws_a.cell(2, 8).value == "À appeler"
    assert ws_a.cell(2, 7).value == "15.01.2026"


def test_export_liste_filtree_single_sheet_and_columns():
    rows = [
        _row(id="1", conseiller="Agent A"),
        _row(id="2", client_contacte=True, conseiller="Agent B", statut_suivi="Signé"),
    ]
    docs_map = {"1": [{"kind": "courrier_optimisation_fiscale", "created_at": "2026-01-15T10:00:00"}]}
    data = build_suivi_3p_export_xlsx(
        rows,
        docs_map,
        scope="liste_filtree",
        grouping="feuille_unique",
        columns=["client", "telephone", "statut"],
    )
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["Export"]
    ws = wb["Export"]
    assert [ws.cell(1, i).value for i in range(1, 4)] == [
        "Nom et prénom du client",
        "Téléphone",
        "Statut du suivi",
    ]
    assert ws.max_row == 3


def test_normalize_export_columns_fallback():
    assert normalize_export_columns(["telephone", "unknown", "email"]) == ["telephone", "email"]
    assert normalize_export_columns([]) == normalize_export_columns(None)
    assert "gain_fiscal" in normalize_export_columns(None)
